"""Stage L — academic DOCX polish for the thesis supervisor draft.

The script assumes Pandoc has already converted the locked thesis Markdown to
DOCX. It then applies deterministic Word styling plus a small OOXML pass for
the parts python-docx cannot express cleanly: a ceremonial title page, section
breaks, page-number format changes, and field updates.
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
XML_NS = "http://www.w3.org/XML/1998/namespace"
NS = {"w": W_NS, "r": R_NS}

FONT = "Times New Roman"
TITLE = "An AI-Based System for Personalized Career Recommendation"
TITLE_PL = "System oparty na sztucznej inteligencji do spersonalizowanego doradztwa zawodowego"
AUTHOR = "Egemen Goncu"
SUPERVISOR = "dr inż. Michał Błędowski"


def set_run_font(run, *, size_pt: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def set_style_font(style, *, size_pt: float, bold: bool | None = None, italic: bool | None = None) -> None:
    style.font.name = FONT
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "666666", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(table, margin_twips: int = 85) -> None:
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def apply_python_docx_polish(src: Path, staged: Path) -> None:
    doc = Document(str(src))

    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    set_style_font(doc.styles["Normal"], size_pt=12)
    normal_pf = doc.styles["Normal"].paragraph_format
    normal_pf.line_spacing = 1.5
    normal_pf.space_after = Pt(6)
    normal_pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size in (("Body Text", 12), ("First Paragraph", 12)):
        if name in doc.styles:
            set_style_font(doc.styles[name], size_pt=size)
            pf = doc.styles[name].paragraph_format
            pf.line_spacing = 1.5
            pf.space_after = Pt(6)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        if name in doc.styles:
            set_style_font(doc.styles[name], size_pt=size, bold=True, italic=(name == "Heading 3"))
            pf = doc.styles[name].paragraph_format
            pf.line_spacing = 1.15
            pf.space_before = Pt(10 if name == "Heading 1" else 8)
            pf.space_after = Pt(8 if name == "Heading 1" else 4)
            pf.keep_with_next = True

    previous_was_heading = False
    for p in doc.paragraphs:
        text = p.text.strip()
        style_name = p.style.name
        if style_name.startswith("Heading"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            if style_name == "Heading 1":
                p.paragraph_format.page_break_before = True
            for run in p.runs:
                set_run_font(run, size_pt=18 if style_name == "Heading 1" else 14 if style_name == "Heading 2" else 12, bold=True, italic=(style_name == "Heading 3"))
            previous_was_heading = True
            continue

        if p._p.findall(".//" + qn("w:drawing")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
            previous_was_heading = False
            continue

        if text:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            if not previous_was_heading and not text.startswith(("Figure ", "Table ", "[", "Keywords:", "Słowa kluczowe:")):
                p.paragraph_format.first_line_indent = Cm(0.5)
            if text.startswith(("Figure ", "Table ")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_together = True
                p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                set_run_font(run, size_pt=12)
        previous_was_heading = False

    max_width = Cm(15.8)
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = int(shape.height * ratio)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_cell_margins(table, 110)
        for r_idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_borders(cell)
                if r_idx == 0:
                    set_cell_shading(cell, "EDEDED")
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        set_run_font(run, size_pt=10.5, bold=(r_idx == 0))

    doc.save(str(staged))


def w(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


def r(tag: str) -> str:
    return f"{{{R_NS}}}{tag}"


def wx(tag: str, **attrs) -> etree._Element:
    element = etree.Element(w(tag))
    for key, value in attrs.items():
        element.set(w(key), str(value))
    return element


def paragraph(
    text: str = "",
    *,
    align: str = "left",
    size: int = 24,
    bold: bool = False,
    italic: bool = False,
    before: int = 0,
    after: int = 0,
    line: int | None = None,
    page_break_before: bool = False,
    style: str | None = None,
) -> etree._Element:
    p = etree.Element(w("p"))
    ppr = etree.SubElement(p, w("pPr"))
    if style:
        pstyle = etree.SubElement(ppr, w("pStyle"))
        pstyle.set(w("val"), style)
    jc = etree.SubElement(ppr, w("jc"))
    jc.set(w("val"), align)
    spacing = etree.SubElement(ppr, w("spacing"))
    spacing.set(w("before"), str(before))
    spacing.set(w("after"), str(after))
    if line:
        spacing.set(w("line"), str(line))
        spacing.set(w("lineRule"), "auto")
    if page_break_before:
        etree.SubElement(ppr, w("pageBreakBefore"))
    if not text:
        return p

    run = etree.SubElement(p, w("r"))
    rpr = etree.SubElement(run, w("rPr"))
    rfonts = etree.SubElement(rpr, w("rFonts"))
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(w(key), FONT)
    sz = etree.SubElement(rpr, w("sz"))
    sz.set(w("val"), str(size))
    szcs = etree.SubElement(rpr, w("szCs"))
    szcs.set(w("val"), str(size))
    if bold:
        etree.SubElement(rpr, w("b"))
        etree.SubElement(rpr, w("bCs"))
    if italic:
        etree.SubElement(rpr, w("i"))
        etree.SubElement(rpr, w("iCs"))
    t = etree.SubElement(run, w("t"))
    if text.startswith(" ") or text.endswith(" "):
        t.set(f"{{{XML_NS}}}space", "preserve")
    t.text = text
    return p


def title_page_paragraphs() -> list[etree._Element]:
    paras = [
        paragraph("Wrocław University of Science and Technology", align="center", size=28, bold=True, after=80),
        paragraph("Faculty of Electronics, Photonics and Microsystems", align="center", size=26, after=620),
        paragraph("Field of Study: Electronic and Computer Engineering", align="center", size=24, after=520),
        paragraph("BACHELOR THESIS", align="center", size=36, bold=True, after=260),
        paragraph("Title of Thesis:", align="center", size=24, after=120),
        paragraph(TITLE, align="center", size=32, bold=True, after=120),
        paragraph(TITLE_PL, align="center", size=24, italic=True, after=620),
        paragraph("Author:", align="center", size=24, after=80),
        paragraph(AUTHOR, align="center", size=24, bold=True, after=260),
        paragraph("Supervisor:", align="center", size=24, after=80),
        paragraph(SUPERVISOR, align="center", size=24, after=780),
        paragraph("WROCŁAW 2026", align="center", size=24, bold=True, after=0),
    ]
    section_p = paragraph("")
    br_run = etree.SubElement(section_p, w("r"))
    br = etree.SubElement(br_run, w("br"))
    br.set(w("type"), "page")
    section_p.find(w("pPr")).append(section_properties(page_number_format=None, footer_rid=None))
    paras.append(section_p)
    return paras


def section_properties(*, page_number_format: str | None, footer_rid: str | None) -> etree._Element:
    sect = etree.Element(w("sectPr"))
    sect_type = etree.SubElement(sect, w("type"))
    sect_type.set(w("val"), "nextPage")
    if footer_rid:
        footer = etree.SubElement(sect, w("footerReference"))
        footer.set(w("type"), "default")
        footer.set(r("id"), footer_rid)
    pg_sz = etree.SubElement(sect, w("pgSz"))
    pg_sz.set(w("w"), "11906")
    pg_sz.set(w("h"), "16838")
    pg_mar = etree.SubElement(sect, w("pgMar"))
    pg_mar.set(w("top"), "1417")
    pg_mar.set(w("right"), "1417")
    pg_mar.set(w("bottom"), "1417")
    pg_mar.set(w("left"), "1417")
    pg_mar.set(w("header"), "709")
    pg_mar.set(w("footer"), "709")
    pg_mar.set(w("gutter"), "0")
    if page_number_format:
        pg_num = etree.SubElement(sect, w("pgNumType"))
        pg_num.set(w("fmt"), page_number_format)
        pg_num.set(w("start"), "1")
    cols = etree.SubElement(sect, w("cols"))
    cols.set(w("space"), "708")
    doc_grid = etree.SubElement(sect, w("docGrid"))
    doc_grid.set(w("linePitch"), "360")
    return sect


def paragraph_text(p: etree._Element) -> str:
    return "".join(t.text or "" for t in p.xpath(".//w:t", namespaces=NS)).strip()


def toc_paragraphs(footer_rid: str) -> list[etree._Element]:
    title = paragraph("Contents", align="left", size=36, bold=True, before=0, after=160, page_break_before=True)
    p = etree.Element(w("p"))
    ppr = etree.SubElement(p, w("pPr"))
    etree.SubElement(ppr, w("keepNext"))
    spacing = etree.SubElement(ppr, w("spacing"))
    spacing.set(w("before"), "0")
    spacing.set(w("after"), "160")
    fld_begin = etree.SubElement(etree.SubElement(p, w("r")), w("fldChar"))
    fld_begin.set(w("fldCharType"), "begin")
    fld_begin.set(w("dirty"), "true")
    instr_run = etree.SubElement(p, w("r"))
    instr = etree.SubElement(instr_run, w("instrText"))
    instr.set(f"{{{XML_NS}}}space", "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = etree.SubElement(etree.SubElement(p, w("r")), w("fldChar"))
    fld_sep.set(w("fldCharType"), "separate")
    placeholder = etree.SubElement(etree.SubElement(p, w("r")), w("t"))
    placeholder.text = "Right-click and update field to refresh the table of contents."
    fld_end = etree.SubElement(etree.SubElement(p, w("r")), w("fldChar"))
    fld_end.set(w("fldCharType"), "end")
    ppr.append(section_properties(page_number_format="lowerRoman", footer_rid=footer_rid))
    return [title, p]


def footer_xml() -> bytes:
    ftr = etree.Element(w("ftr"), nsmap={"w": W_NS, "r": R_NS})
    p = etree.SubElement(ftr, w("p"))
    ppr = etree.SubElement(p, w("pPr"))
    jc = etree.SubElement(ppr, w("jc"))
    jc.set(w("val"), "center")
    for tag, text in (
        ("begin", None),
        ("instr", " PAGE "),
        ("separate", None),
        ("result", "1"),
        ("end", None),
    ):
        run = etree.SubElement(p, w("r"))
        if tag == "instr":
            instr = etree.SubElement(run, w("instrText"))
            instr.set(f"{{{XML_NS}}}space", "preserve")
            instr.text = text
        elif tag == "result":
            t = etree.SubElement(run, w("t"))
            t.text = text
        else:
            fld = etree.SubElement(run, w("fldChar"))
            fld.set(w("fldCharType"), tag)
    return etree.tostring(ftr, xml_declaration=True, encoding="UTF-8", standalone="yes")


def next_footer_name(names: set[str]) -> str:
    i = 1
    while f"word/footer{i}.xml" in names:
        i += 1
    return f"footer{i}.xml"


def add_footer_relationship(rels_root: etree._Element, footer_name: str) -> str:
    max_id = 0
    for rel in rels_root.findall(f"{{{REL_NS}}}Relationship"):
        rid = rel.get("Id", "")
        if rid.startswith("rId") and rid[3:].isdigit():
            max_id = max(max_id, int(rid[3:]))
    rid = f"rId{max_id + 1}"
    rel = etree.SubElement(rels_root, f"{{{REL_NS}}}Relationship")
    rel.set("Id", rid)
    rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer")
    rel.set("Target", footer_name)
    return rid


def add_content_type_override(ct_root: etree._Element, footer_name: str) -> None:
    part = f"/word/{footer_name}"
    for node in ct_root.findall(f"{{{CT_NS}}}Override"):
        if node.get("PartName") == part:
            return
    override = etree.SubElement(ct_root, f"{{{CT_NS}}}Override")
    override.set("PartName", part)
    override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml")


def mark_update_fields(settings_root: etree._Element) -> None:
    node = settings_root.find("w:updateFields", namespaces=NS)
    if node is None:
        node = etree.SubElement(settings_root, w("updateFields"))
    node.set(w("val"), "true")


def apply_ooxml_polish(staged: Path, dst: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        with zipfile.ZipFile(staged, "r") as zin:
            zin.extractall(tmp_path)

        document_path = tmp_path / "word" / "document.xml"
        rels_path = tmp_path / "word" / "_rels" / "document.xml.rels"
        ct_path = tmp_path / "[Content_Types].xml"
        settings_path = tmp_path / "word" / "settings.xml"

        parser = etree.XMLParser(remove_blank_text=False)
        doc_root = etree.parse(str(document_path), parser).getroot()
        rels_root = etree.parse(str(rels_path), parser).getroot()
        ct_root = etree.parse(str(ct_path), parser).getroot()
        settings_root = etree.parse(str(settings_path), parser).getroot()
        body = doc_root.find("w:body", namespaces=NS)
        if body is None:
            raise RuntimeError("document.xml has no w:body")

        names = {str(p.relative_to(tmp_path)) for p in tmp_path.rglob("*") if p.is_file()}
        footer_name = next_footer_name(names)
        footer_rid = add_footer_relationship(rels_root, footer_name)
        add_content_type_override(ct_root, footer_name)
        mark_update_fields(settings_root)

        final_sect = body.find("w:sectPr", namespaces=NS)
        if final_sect is not None:
            body.remove(final_sect)

        abstract_idx = None
        for i, child in enumerate(body):
            if child.tag == w("p") and paragraph_text(child).startswith("Abstract"):
                abstract_idx = i
                break
        if abstract_idx is None:
            raise RuntimeError("Could not find Abstract heading after Pandoc conversion")
        for child in list(body)[:abstract_idx]:
            body.remove(child)

        abstract_p = body[0]
        abstract_ppr = abstract_p.find(w("pPr"))
        if abstract_ppr is None:
            abstract_ppr = etree.Element(w("pPr"))
            abstract_p.insert(0, abstract_ppr)
        if abstract_ppr.find(w("pageBreakBefore")) is None:
            etree.SubElement(abstract_ppr, w("pageBreakBefore"))

        first = body[0]
        for p in title_page_paragraphs():
            body.insert(body.index(first), p)

        chapter_idx = None
        for i, child in enumerate(body):
            if child.tag == w("p") and paragraph_text(child).startswith("Chapter 1"):
                chapter_idx = i
                break
        if chapter_idx is None:
            raise RuntimeError("Could not find Chapter 1 heading for body section break")
        for p in toc_paragraphs(footer_rid):
            body.insert(chapter_idx, p)

        body.append(section_properties(page_number_format="decimal", footer_rid=footer_rid))

        (tmp_path / "word" / footer_name).write_bytes(footer_xml())
        document_path.write_bytes(etree.tostring(doc_root, xml_declaration=True, encoding="UTF-8", standalone="yes"))
        rels_path.write_bytes(etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes"))
        ct_path.write_bytes(etree.tostring(ct_root, xml_declaration=True, encoding="UTF-8", standalone="yes"))
        settings_path.write_bytes(etree.tostring(settings_root, xml_declaration=True, encoding="UTF-8", standalone="yes"))

        dst.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for path in sorted(tmp_path.rglob("*")):
                if path.is_file():
                    zout.write(path, path.relative_to(tmp_path))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--in", dest="src", required=True)
    parser.add_argument("--out", dest="dst", required=True)
    args = parser.parse_args()

    src = Path(args.src)
    dst = Path(args.dst)
    if not src.exists():
        print(f"ERROR: input DOCX not found: {src}", file=sys.stderr)
        return 1

    with tempfile.TemporaryDirectory() as tmp:
        staged = Path(tmp) / "stage_l_python_docx.docx"
        apply_python_docx_polish(src, staged)
        apply_ooxml_polish(staged, dst)

    print(f"Saved polished thesis DOCX: {dst}")
    print("Applied: title page, Times 12 body, 1.5 spacing, justified text, margins, table styling, figures, TOC field, Roman/Arabic section numbering.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
