"""B6: CV render, export, ATS checks, versions and concurrent edits, at the HTTP seam.

Expected values come from the product contract (a long CV flows across pages
without a blank first page or a split role; text the font cannot draw is reported,
not dropped silently; a second writer is told, not overwritten), not from the code.
"""

import io

import fitz
import pytest
from docx import Document
from docx.shared import Mm

from app.schemas.cv_documents import CvDocumentCreate, CvStyle
from app.services.cv_documents import create_document
from app.services.cv_html import load_manifest
from app.services.cv_pdf import embedded_fonts, font_problems
from app.services.cv_rendering import build_render_model, validate_artifact
from tests.test_cv_documents import _signed

PREFIX = "/api/v1/cv-documents"
A4_HEIGHT_PT = 842


def _entry(entry_id, position, **fields):
    entry = {"id": entry_id, "evidence_item_id": None, "position": position}
    entry.update(fields)
    entry.setdefault("body", fields.get("heading") or "Text")
    return entry


def _section(kind, title, position, entries):
    return {
        "id": f"s-{kind}",
        "kind": kind,
        "title": title,
        "visible": True,
        "position": position,
        "entries": entries,
    }


def _role(index, bullet_count=5):
    return _entry(
        f"role-{index}",
        index,
        heading=f"Role {index} Engineer",
        subheading=f"Company {index}",
        start_date="2016",
        end_date="2020",
        bullets=[
            f"Role {index} bullet {n}: delivered a measurable improvement to a production "
            f"system used by several teams and customers."
            for n in range(1, bullet_count + 1)
        ],
    )


def mid_cv_sections():
    """Summary, eight roles of five bullets, Skills and Education (the audit's mid2 fixture)."""
    return [
        _section(
            "summary",
            "Summary",
            0,
            [_entry("sum", 0, body="Backend engineer with a decade of experience building services.")],
        ),
        _section("experience", "Experience", 1, [_role(i) for i in range(1, 9)]),
        _section("skills", "Skills", 2, [_entry("sk", 0, body="Python, Go, PostgreSQL, Kubernetes")]),
        _section(
            "education",
            "Education",
            3,
            [
                _entry(
                    "ed",
                    0,
                    heading="BSc Computer Science",
                    subheading="University of Berlin",
                    start_date="2008",
                    end_date="2012",
                )
            ],
        ),
    ]


def _create(client, auth_headers, sections, name="Test CV", header=None):
    body = {"name": name, "sections": sections}
    if header is not None:
        body["header"] = header
    response = client.post(PREFIX, json=body, headers=auth_headers)
    assert response.status_code == 201, response.text
    return response.json()


def _pdf(client, auth_headers, document_id, variant=None):
    path = f"variants/{variant}/artifacts/pdf" if variant else "artifacts/pdf"
    response = client.get(f"{PREFIX}/{document_id}/{path}", headers=auth_headers)
    assert response.status_code == 200, response.text
    return response.content


def _quality(client, auth_headers, document_id, **params):
    response = client.post(
        f"{PREFIX}/{document_id}/quality", params=params, headers=auth_headers
    )
    assert response.status_code == 200, response.text
    return {check["id"]: check for check in response.json()["checks"]}


# ── d01: pagination ──


def test_long_cv_fills_the_first_page_and_never_splits_a_role(client, auth_headers):
    document = _create(client, auth_headers, mid_cv_sections())
    pdf = _pdf(client, auth_headers, document["id"])

    with fitz.open(stream=pdf, filetype="pdf") as parsed:
        pages = [page.get_text() for page in parsed]
        # The first page carries real content, not just the title and Summary.
        first_page_bottom = max(block[3] for block in parsed[0].get_text("blocks"))
        assert first_page_bottom > 0.7 * A4_HEIGHT_PT
        assert "Role 1 Engineer" in pages[0]
        # Every role is whole: its heading and all five bullets share one page.
        for index in range(1, 9):
            holders = [i for i, text in enumerate(pages) if f"Role {index} Engineer" in text]
            assert len(holders) == 1
            text = pages[holders[0]]
            assert all(f"Role {index} bullet {n}:" in text for n in range(1, 6))
        # Nothing was lost to the reflow.
        assert "Education" in "".join(pages) and "BSc Computer Science" in "".join(pages)


def test_long_cv_passes_every_ats_check(client, auth_headers):
    document = _create(client, auth_headers, mid_cv_sections())
    checks = _quality(client, auth_headers, document["id"])
    assert all(check["passed"] for check in checks.values()), checks


def _page_gap_pdf(page_one_lines, page_two_lines):
    pdf = fitz.open()
    for lines in (page_one_lines, page_two_lines):
        page = pdf.new_page(width=595, height=842)
        y = 60
        for line in lines:
            page.insert_text((50, y), line, fontsize=11)
            y += 16
    data = pdf.tobytes()
    pdf.close()
    return data


def test_a_large_gap_at_the_bottom_of_a_non_final_page_fails_tidy_page_breaks(db, test_user):
    document = create_document(
        db,
        test_user.id,
        CvDocumentCreate(
            name="Gap", sections=[_section("summary", "Summary", 0, [_entry("a", 0, body="One short line.")])]
        ),
    )
    model = build_render_model(document, "classic", CvStyle())

    # Page one holds a title and the summary only; the rest sits on page two.
    sparse = _page_gap_pdf(["Test User", "Summary", "One short line."], ["More content on page two."])
    assert validate_artifact(model, sparse).page_breaks == "fail"
    # A page that is simply full of content is fine.
    full = _page_gap_pdf(["Test User", "Summary", "One short line."] + ["filler"] * 44, ["More."])
    assert validate_artifact(model, full).page_breaks == "pass"


def test_legacy_two_column_template_prints_as_the_single_column_classic_for_now(
    client, auth_headers
):
    # TODO(T2/T6): the sidebar templates return with T6; until then the legacy
    # two-column id prints as `classic`, so no content is lost or moved into a sidebar.
    document = _create(client, auth_headers, mid_cv_sections())
    client.patch(
        f"{PREFIX}/{document['id']}",
        json={"style": {"template_id": "lagoon"}},
        headers=auth_headers,
    )
    with fitz.open(stream=_pdf(client, auth_headers, document["id"]), filetype="pdf") as parsed:
        assert parsed.page_count > 1
        # Single column: page two's text starts at the page margin, not right of a sidebar.
        assert min(block[0] for block in parsed[1].get_text("blocks")) < 58 * 72 / 25.4


# ── d10: ATS check, imported CV ──

IMPORTED = """Casey Morgan
casey@example.com | Berlin, Germany

Summary
• Platform engineer focused on reliable systems.
- Mentors juniors.

Professional Experience
Senior Engineer at Northwind (2021 - Present)
* Cut p95 latency by 38%.
– Led 4 engineers through a migration.

Skills
• Python, Go
- Docker
"""


def test_an_imported_single_column_cv_reads_back(client, auth_headers):
    proposal = client.post(
        f"{PREFIX}/import/proposals",
        files={"file": ("cv.txt", io.BytesIO(IMPORTED.encode()), "text/plain")},
        headers=auth_headers,
    ).json()
    document = client.post(f"{PREFIX}/import/accept", json=proposal, headers=auth_headers).json()
    checks = _quality(client, auth_headers, document["id"])
    assert checks["reads_back"]["passed"], checks["reads_back"]
    assert checks["page_breaks"]["passed"]


# ── d08: text the font cannot draw ──


def _body_cv(text):
    return [_section("summary", "Summary", 0, [_entry("a", 0, body=text)])]


def _pdf_text(pdf):
    with fitz.open(stream=pdf, filetype="pdf") as parsed:
        return " ".join(page.get_text() for page in parsed)


@pytest.mark.parametrize(
    "name", ["Дмитрий Соколов", "Αλέξανδρος Παπαδόπουλος", "Nguyễn Thị Hương", "Łukasz Żółć"]
)
@pytest.mark.parametrize(
    "style",
    [
        {"font_id": "lato"},
        {"font_id": "crimson-text", "template_id": "executive"},
        {"font_id": "ibm-plex-mono", "template_id": "slate"},
        {"font_id": "pt-sans", "template_id": "executive"},
        {"ats_mode": True},
    ],
    ids=["lato", "crimson", "plex-mono", "pt-sans", "ats-mode"],
)
def test_names_in_scripts_some_bundled_font_covers_survive_the_pdf(
    client, auth_headers, name, style
):
    document = _create(client, auth_headers, _body_cv("Engineer."), header={"name": name})
    client.patch(f"{PREFIX}/{document['id']}", json={"style": style}, headers=auth_headers)
    assert name in _pdf_text(_pdf(client, auth_headers, document["id"]))
    assert _quality(client, auth_headers, document["id"])["reads_back"]["passed"]


def test_characters_no_bundled_font_can_draw_are_reported_not_silently_dropped(
    client, auth_headers
):
    document = _create(client, auth_headers, _body_cv("日本語テスト 😀 Backend engineer."))
    check = _quality(client, auth_headers, document["id"])["reads_back"]
    assert check["passed"] is False
    assert "日" in check["fix"] and "😀" in check["fix"]
    assert "ATS-friendly mode" not in check["fix"]


# ── d24: the reads-back advice names what failed ──


def test_reads_back_failure_names_the_section_that_did_not_read_back(client, auth_headers):
    sections = [
        _section("summary", "Summary", 0, [_entry("a", 0, body="Fine summary.")]),
        _section("experience", "Experience", 1, [_entry("b", 0, body="日本語テスト Backend.")]),
        _section("skills", "Skills", 2, [_entry("c", 0, body="Python")]),
    ]
    document = _create(client, auth_headers, sections)
    check = _quality(client, auth_headers, document["id"])["reads_back"]
    assert check["passed"] is False
    assert "Experience" in check["fix"]
    assert "Skills" not in check["fix"]


# ── d23: section headings check ──


@pytest.mark.parametrize(
    "kinds, expected",
    [
        (("experience", "skills"), True),
        (("education", "projects", "skills"), True),  # a graduate CV needs no Experience
        (("summary", "skills"), False),  # nothing says what the person did or studied
        (("experience",), False),  # a single heading is not a structure
    ],
)
def test_section_headings_check_accepts_standard_structures(client, auth_headers, kinds, expected):
    titles = {"experience": "Experience", "skills": "Skills", "education": "Education",
              "projects": "Projects", "summary": "Summary"}
    sections = [
        _section(kind, titles[kind], index, [_entry(f"e{index}", 0, body=f"{titles[kind]} text")])
        for index, kind in enumerate(kinds)
    ]
    document = _create(client, auth_headers, sections)
    assert _quality(client, auth_headers, document["id"])["sections"]["passed"] is expected


# ── d22: embedded fonts ──


def test_pdf_embeds_every_font_it_uses(client, auth_headers):
    document = _create(client, auth_headers, mid_cv_sections())
    client.patch(
        f"{PREFIX}/{document['id']}",
        json={"style": {"template_id": "executive", "font_id": "pt-serif"}},
        headers=auth_headers,
    )
    pdf = _pdf(client, auth_headers, document["id"])
    with fitz.open(stream=pdf, filetype="pdf") as parsed:
        fonts = {font[3]: font[1] for page in parsed for font in page.get_fonts()}
    assert fonts
    assert all(kind != "n/a" for kind in fonts.values()), fonts


# ── d03: two writers ──


def test_a_second_writer_with_a_stale_copy_gets_409_with_the_newer_version(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Original."))
    first_load = document["updated_at"]
    sections = document["sections"]

    tab_one = client.patch(
        f"{PREFIX}/{document['id']}",
        json={"name": "Renamed in tab one", "expected_updated_at": first_load},
        headers=auth_headers,
    )
    assert tab_one.status_code == 200

    sections[0]["entries"][0]["body"] = "Edited in tab two."
    tab_two = client.patch(
        f"{PREFIX}/{document['id']}",
        json={"sections": sections, "expected_updated_at": first_load},
        headers=auth_headers,
    )
    assert tab_two.status_code == 409
    body = tab_two.json()
    assert isinstance(body["detail"], str) and body["detail"]
    assert body["current"]["name"] == "Renamed in tab one"
    assert body["current"]["updated_at"] == tab_one.json()["updated_at"]
    # The loser did not overwrite anything.
    stored = client.get(f"{PREFIX}/{document['id']}", headers=auth_headers).json()
    assert stored["sections"][0]["entries"][0]["body"] == "Original."

    # Retrying on the newer version succeeds; omitting the field stays allowed.
    retry = client.patch(
        f"{PREFIX}/{document['id']}",
        json={"sections": sections, "expected_updated_at": body["current"]["updated_at"]},
        headers=auth_headers,
    )
    assert retry.status_code == 200
    unconditional = client.patch(
        f"{PREFIX}/{document['id']}", json={"name": "No precondition"}, headers=auth_headers
    )
    assert unconditional.status_code == 200


def test_counting_a_tailoring_run_does_not_invalidate_an_open_editor(
    client, auth_headers, monkeypatch
):
    async def no_changes(*_args, **_kwargs):
        return {"changes": []}

    monkeypatch.setattr("app.services.cv_tailoring.complete_structured", no_changes)
    document = _create(client, auth_headers, _body_cv("Original."))
    proposal = client.post(
        f"{PREFIX}/{document['id']}/tailoring",
        json={
            "job_title": "Platform Engineer",
            "job_description": "Platform engineer for distributed backend services.",
        },
        headers=auth_headers,
    )
    assert proposal.status_code == 200, proposal.text
    save = client.patch(
        f"{PREFIX}/{document['id']}",
        json={"name": "Still mine", "expected_updated_at": document["updated_at"]},
        headers=auth_headers,
    )
    assert save.status_code == 200, save.text


# ── d11: a saved version can be previewed, exported, renamed and deleted ──


def test_a_version_exports_without_touching_the_working_cv(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Working copy text."))
    variant = client.post(
        f"{PREFIX}/{document['id']}/variants",
        json={"name": "For Acme", "target_role": "Platform Engineer"},
        headers=auth_headers,
    ).json()
    sections = document["sections"]
    sections[0]["entries"][0]["body"] = "Edited after saving the version."
    client.patch(f"{PREFIX}/{document['id']}", json={"sections": sections}, headers=auth_headers)

    version_pdf = _pdf(client, auth_headers, document["id"], variant=variant["id"])
    working_pdf = _pdf(client, auth_headers, document["id"])
    assert "Working copy text." in _pdf_text(version_pdf)
    assert "Edited after saving" not in _pdf_text(version_pdf)
    assert "Edited after saving" in _pdf_text(working_pdf)

    docx = client.get(
        f"{PREFIX}/{document['id']}/variants/{variant['id']}/artifacts/docx", headers=auth_headers
    )
    assert docx.status_code == 200
    assert "Working copy text." in "\n".join(
        p.text for p in Document(io.BytesIO(docx.content)).paragraphs
    )
    checks = _quality(client, auth_headers, document["id"], variant_id=variant["id"])
    assert checks["reads_back"]["passed"]
    # The working CV is exactly as the person left it.
    current = client.get(f"{PREFIX}/{document['id']}", headers=auth_headers).json()
    assert current["sections"][0]["entries"][0]["body"] == "Edited after saving the version."


def test_an_unknown_or_foreign_version_id_is_a_404(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."))
    response = client.get(
        f"{PREFIX}/{document['id']}/variants/not-a-version/artifacts/pdf", headers=auth_headers
    )
    assert response.status_code == 404


def test_versions_can_be_renamed_and_deleted(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."))
    first = client.post(
        f"{PREFIX}/{document['id']}/variants", json={"name": "One"}, headers=auth_headers
    ).json()
    client.post(f"{PREFIX}/{document['id']}/variants", json={"name": "Two"}, headers=auth_headers)

    renamed = client.patch(
        f"{PREFIX}/{document['id']}/variants/{first['id']}",
        json={"name": "For Acme", "target_role": "Staff Engineer"},
        headers=auth_headers,
    )
    assert renamed.status_code == 200
    assert renamed.json()["name"] == "For Acme" and renamed.json()["target_role"] == "Staff Engineer"
    clash = client.patch(
        f"{PREFIX}/{document['id']}/variants/{first['id']}",
        json={"name": "Two"},
        headers=auth_headers,
    )
    assert clash.status_code == 409

    deleted = client.delete(
        f"{PREFIX}/{document['id']}/variants/{first['id']}", headers=auth_headers
    )
    assert deleted.status_code == 204
    names = [
        v["name"] for v in client.get(f"{PREFIX}/{document['id']}", headers=auth_headers).json()["variants"]
    ]
    assert "For Acme" not in names and "Two" in names
    assert (
        client.delete(
            f"{PREFIX}/{document['id']}/variants/{first['id']}", headers=auth_headers
        ).status_code
        == 404
    )


def test_deleting_a_version_clears_it_from_applications_that_selected_it(
    client, auth_headers, db, test_user
):
    from app.models.workspace import Workspace  # noqa: PLC0415

    document = _create(client, auth_headers, _body_cv("Text."))
    variant = client.post(
        f"{PREFIX}/{document['id']}/variants", json={"name": "For Acme"}, headers=auth_headers
    ).json()
    workspace = Workspace(user_id=test_user.id, label="Acme", selected_cv_variant_id=variant["id"])
    db.add(workspace)
    db.commit()
    assert (
        client.delete(
            f"{PREFIX}/{document['id']}/variants/{variant['id']}", headers=auth_headers
        ).status_code
        == 204
    )
    db.refresh(workspace)
    assert workspace.selected_cv_variant_id is None


# ── d15 / d16: input validation ──


def test_nul_characters_are_rejected_as_422_not_500(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."))
    response = client.patch(
        f"{PREFIX}/{document['id']}", json={"name": "Bad\u0000name"}, headers=auth_headers
    )
    assert response.status_code == 422
    sections = _body_cv("Bad\u0000body")
    assert (
        client.post(PREFIX, json={"name": "x", "sections": sections}, headers=auth_headers).status_code
        == 422
    )


@pytest.mark.parametrize("blank", ["", "   ", "\n\t"])
def test_blank_names_titles_and_bodies_are_rejected(client, auth_headers, blank):
    document = _create(client, auth_headers, _body_cv("Text."))
    assert (
        client.patch(
            f"{PREFIX}/{document['id']}", json={"name": blank}, headers=auth_headers
        ).status_code
        == 422
    )
    assert (
        client.post(
            f"{PREFIX}/{document['id']}/variants", json={"name": blank}, headers=auth_headers
        ).status_code
        == 422
    )
    for mutate in (
        lambda s: s[0].update(title=blank),
        lambda s: s[0]["entries"][0].update(body=blank),
    ):
        sections = _body_cv("Text.")
        mutate(sections)
        assert (
            client.patch(
                f"{PREFIX}/{document['id']}", json={"sections": sections}, headers=auth_headers
            ).status_code
            == 422
        )


def test_names_are_stored_trimmed(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."), name="  Padded name  ")
    assert document["name"] == "Padded name"


def test_duplicate_section_and_entry_ids_are_rejected(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."))
    twin_sections = [
        _section("summary", "Summary", 0, [_entry("a", 0, body="x")]),
        _section("summary", "Summary again", 1, [_entry("b", 0, body="y")]),
    ]
    assert (
        client.patch(
            f"{PREFIX}/{document['id']}", json={"sections": twin_sections}, headers=auth_headers
        ).status_code
        == 422
    )
    twin_entries = [_section("summary", "Summary", 0, [_entry("a", 0, body="x"), _entry("a", 1, body="y")])]
    assert (
        client.patch(
            f"{PREFIX}/{document['id']}", json={"sections": twin_entries}, headers=auth_headers
        ).status_code
        == 422
    )


def test_a_second_user_cannot_export_someone_elses_version(client, auth_headers, db):
    from app.auth.security import create_access_token, hash_password  # noqa: PLC0415
    from app.models.user import User  # noqa: PLC0415

    document = _create(client, auth_headers, _body_cv("Private."))
    variant = client.post(
        f"{PREFIX}/{document['id']}/variants", json={"name": "Mine"}, headers=auth_headers
    ).json()
    other = User(email="intruder@example.com", hashed_password=hash_password("password123"))
    db.add(other)
    db.commit()
    headers = {"Authorization": f"Bearer {create_access_token(other.id)}"}
    url = f"{PREFIX}/{document['id']}/variants/{variant['id']}/artifacts/pdf"
    assert client.get(url, headers=headers).status_code == 404


# ── d17: separate export limits ──


def test_pdf_downloads_do_not_use_up_the_docx_limit(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."))
    for _ in range(10):
        assert client.get(f"{PREFIX}/{document['id']}/artifacts/pdf", headers=auth_headers).status_code == 200
    assert client.get(f"{PREFIX}/{document['id']}/artifacts/pdf", headers=auth_headers).status_code == 429
    assert client.get(f"{PREFIX}/{document['id']}/artifacts/docx", headers=auth_headers).status_code == 200


# ── d21: replays are 200 ──


def test_replaying_an_import_accept_returns_200_with_the_same_document(client, auth_headers):
    proposal = client.post(
        f"{PREFIX}/import/proposals",
        files={"file": ("cv.txt", io.BytesIO(IMPORTED.encode()), "text/plain")},
        headers=auth_headers,
    ).json()
    first = client.post(f"{PREFIX}/import/accept", json=proposal, headers=auth_headers)
    replay = client.post(f"{PREFIX}/import/accept", json=proposal, headers=auth_headers)
    assert first.status_code == 201
    assert replay.status_code == 200 and replay.json()["id"] == first.json()["id"]


def test_replaying_a_tailoring_apply_returns_200_with_the_same_version(
    client, auth_headers, test_user
):
    document = _create(client, auth_headers, _body_cv("Platform work."))
    change = {
        "id": "c1",
        "section_id": "s-summary",
        "entry_id": "a",
        "before": "Platform work.",
        "after": "Reliable platform work.",
        "job_requirement": "Reliability",
        "evidence_item_ids": [],
        "support": "document",
    }
    payload = _signed(
        document["id"],
        test_user.id,
        {
            "request_id": "b5ac39c0-5a76-4e94-98f1-a0fd8b42a5b2",
            "variant_name": "Tailored",
            "job_title": "Platform Engineer",
            "changes": [change],
            "decisions": [{"change_id": "c1", "action": "accept"}],
        },
    )
    first = client.post(f"{PREFIX}/{document['id']}/tailoring/apply", json=payload, headers=auth_headers)
    replay = client.post(f"{PREFIX}/{document['id']}/tailoring/apply", json=payload, headers=auth_headers)
    assert first.status_code == 201
    assert replay.status_code == 200 and replay.json()["id"] == first.json()["id"]


# ── d14: the DOCX is an A4 document with real headings ──


def test_docx_is_a4_with_heading_styles(client, auth_headers):
    document = _create(client, auth_headers, mid_cv_sections())
    response = client.get(f"{PREFIX}/{document['id']}/artifacts/docx", headers=auth_headers)
    parsed = Document(io.BytesIO(response.content))
    section = parsed.sections[0]
    assert abs(section.page_width - Mm(210)) < Mm(1)
    assert abs(section.page_height - Mm(297)) < Mm(1)
    headings = [p.text for p in parsed.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Summary", "Experience", "Skills", "Education"]


# ── review fixes: pagination of one long paragraph ──


def _long_paragraph_sections():
    long_body = " ".join(
        f"Project sentence {n} describes a delivered result for a customer."
        for n in range(1, 60)
    )
    assert len(long_body) > 3_000
    return [
        _section("experience", "Experience", 0, [_role(i, 4) for i in range(1, 5)]),
        _section("projects", "Projects", 1, [_entry("proj", 0, body=long_body)]),
    ]


def test_a_long_single_paragraph_entry_splits_instead_of_leaving_a_blank_page(
    client, auth_headers
):
    document = _create(client, auth_headers, _long_paragraph_sections())
    pdf = _pdf(client, auth_headers, document["id"])
    with fitz.open(stream=pdf, filetype="pdf") as parsed:
        assert parsed.page_count >= 2
        first_page_bottom = max(block[3] for block in parsed[0].get_text("blocks"))
        assert first_page_bottom > 0.6 * A4_HEIGHT_PT
        text = " ".join(page.get_text() for page in parsed)
        # The Projects heading stays with the start of its text on page one.
        assert "Projects" in parsed[0].get_text() and "Project sentence 1 " in parsed[0].get_text()
    flat = " ".join(text.split())
    assert "Project sentence 59 describes" in flat and "Role 4 bullet 4" in flat
    assert _quality(client, auth_headers, document["id"])["page_breaks"]["passed"]


# ── review fixes: stored rows from before blank values were refused ──


def _store_document(db, user, sections, name="Legacy CV"):
    from app.models.cv_document import CvDocument

    document = CvDocument(user_id=user.id, name=name, sections=sections)
    db.add(document)
    db.commit()
    return document


def _legacy_sections():
    return [
        {
            "id": "s1",
            "kind": "experience",
            "title": "Experience",
            "visible": True,
            "position": 0,
            "entries": [
                {
                    "id": "e1",
                    "evidence_item_id": None,
                    "position": 0,
                    "body": "Built things.",
                    "heading": "   ",
                    "subheading": " ",
                    "location": "  ",
                    "start_date": "   ",
                    "end_date": " ",
                    "bullets": ["Real bullet", "   "],
                },
                {"id": "e2", "evidence_item_id": None, "position": 1, "body": "   ", "bullets": []},
            ],
        }
    ]


def test_stored_documents_with_blank_optional_values_still_open_list_export_and_check(
    client, auth_headers, db, test_user
):
    document = _store_document(db, test_user, _legacy_sections())
    fetched = client.get(f"{PREFIX}/{document.id}", headers=auth_headers)
    assert fetched.status_code == 200, fetched.text
    entry = fetched.json()["sections"][0]["entries"][0]
    assert entry["heading"] is None and entry["start_date"] is None
    assert entry["bullets"] == ["Real bullet"]
    listed = client.get(PREFIX, headers=auth_headers).json()["items"]
    assert any(item["id"] == document.id and item["sections"] for item in listed)
    assert client.patch(
        f"{PREFIX}/{document.id}", json={"name": "Still opens"}, headers=auth_headers
    ).status_code == 200
    _pdf(client, auth_headers, document.id)
    assert client.get(f"{PREFIX}/{document.id}/artifacts/docx", headers=auth_headers).status_code == 200
    assert client.post(f"{PREFIX}/{document.id}/quality", headers=auth_headers).status_code == 200


# ── review fixes: the sidebar overflowing page one ──


def test_a_long_sidebar_does_not_push_the_main_column_into_the_sidebar_frame(
    client, auth_headers
):
    skills = [_entry(f"sk{n}", n, body=f"Skill {n} word word word word word") for n in range(60)]
    sections = [
        _section("summary", "Summary", 0, [_entry("sum", 0, body="Backend engineer.")]),
        _section("experience", "Experience", 1, [_role(i) for i in range(1, 3)]),
        _section("skills", "Skills", 2, skills),
    ]
    document = _create(client, auth_headers, sections)
    client.patch(
        f"{PREFIX}/{document['id']}",
        json={"style": {"template_id": "lagoon"}},
        headers=auth_headers,
    )
    with fitz.open(stream=_pdf(client, auth_headers, document["id"]), filetype="pdf") as parsed:
        assert parsed.page_count > 1
        # Pages after the first carry the main column only, never the 58 mm sidebar frame.
        for page in parsed.pages(1):
            for block in page.get_text("blocks"):
                text = block[4]
                if any(f"Role {i} " in text for i in range(1, 3)) or "Experience" in text:
                    assert block[0] >= 58 * 72 / 25.4, (page.number, block)


# ── review fixes: smaller items ──


def test_clearing_selection_for_an_empty_list_of_versions_clears_nothing(db, test_user):
    from app.models.workspace import Workspace
    from app.services.applications import clear_selected_variants

    document = create_document(
        db, test_user.id, CvDocumentCreate(name="Doc", sections=_body_cv("Text."))
    )
    from app.models.cv_document import CvVariant

    variant = CvVariant(document_id=document.id, name="V", sections=document.sections)
    db.add(variant)
    db.commit()
    workspace = Workspace(user_id=test_user.id, label="App", selected_cv_variant_id=variant.id)
    db.add(workspace)
    db.commit()
    clear_selected_variants(db, document, [])
    db.commit()
    db.refresh(workspace)
    assert workspace.selected_cv_variant_id == variant.id


def test_a_symbol_in_the_text_does_not_change_the_whole_cv_font(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Reliable systems \u2713 done."))
    pdf = _pdf(client, auth_headers, document["id"])
    # The template's own families only: no fallback face is substituted for the symbol.
    assert embedded_fonts(pdf)
    assert font_problems(pdf, load_manifest("classic").families) == []


def test_sidebar_template_order_difference_is_explained_as_layout_not_missing_sections(
    client, auth_headers
):
    document = _create(client, auth_headers, mid_cv_sections())
    client.patch(
        f"{PREFIX}/{document['id']}",
        json={"style": {"template_id": "lagoon"}},
        headers=auth_headers,
    )
    check = _quality(client, auth_headers, document["id"])["reads_back"]
    if not check["passed"]:
        assert "single-column" in check["fix"]
        assert "did not read back as written" not in check["fix"]


def test_a_versions_download_name_includes_the_version_name(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."))
    variant = client.post(
        f"{PREFIX}/{document['id']}/variants", json={"name": "For Acme"}, headers=auth_headers
    ).json()
    response = client.get(
        f"{PREFIX}/{document['id']}/variants/{variant['id']}/artifacts/pdf", headers=auth_headers
    )
    assert "For-Acme" in response.headers["content-disposition"]
    working = client.get(f"{PREFIX}/{document['id']}/artifacts/pdf", headers=auth_headers)
    assert "For-Acme" not in working.headers["content-disposition"]


def test_a_cv_too_long_to_read_back_is_reported_not_a_500(client, auth_headers, monkeypatch):
    from app.services import cv_rendering
    from app.services.cv_parser import CvParserRejected

    def too_long(*_args, **_kwargs):
        raise CvParserRejected("too many pages")

    monkeypatch.setattr(cv_rendering, "parse_cv", too_long)
    document = _create(client, auth_headers, _body_cv("Text."))
    checks = _quality(client, auth_headers, document["id"])
    assert checks["reads_back"]["passed"] is False


def test_a_patch_with_only_the_precondition_is_422(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."))
    response = client.patch(
        f"{PREFIX}/{document['id']}",
        json={"expected_updated_at": document["updated_at"]},
        headers=auth_headers,
    )
    assert response.status_code == 422


def test_renaming_a_version_to_its_own_name_is_accepted(client, auth_headers):
    document = _create(client, auth_headers, _body_cv("Text."))
    variant = client.post(
        f"{PREFIX}/{document['id']}/variants", json={"name": "Same"}, headers=auth_headers
    ).json()
    response = client.patch(
        f"{PREFIX}/{document['id']}/variants/{variant['id']}",
        json={"name": "Same", "target_role": "Engineer"},
        headers=auth_headers,
    )
    assert response.status_code == 200 and response.json()["target_role"] == "Engineer"


def test_a_two_column_version_exports(client, auth_headers):
    document = _create(client, auth_headers, mid_cv_sections())
    client.patch(
        f"{PREFIX}/{document['id']}",
        json={"style": {"template_id": "lagoon"}},
        headers=auth_headers,
    )
    variant = client.post(
        f"{PREFIX}/{document['id']}/variants", json={"name": "Two column"}, headers=auth_headers
    ).json()
    assert "Role 1 Engineer" in _pdf_text(_pdf(client, auth_headers, document["id"], variant["id"]))


def test_nul_in_an_import_proposal_is_stripped_not_refused(client, auth_headers):
    proposal = client.post(
        f"{PREFIX}/import/proposals",
        files={"file": ("cv.txt", io.BytesIO(IMPORTED.encode()), "text/plain")},
        headers=auth_headers,
    ).json()
    proposal["name"] = "Casey\x00 Morgan"
    proposal["sections"][0]["entries"][0]["body"] += "\x00 more"
    accepted = client.post(f"{PREFIX}/import/accept", json=proposal, headers=auth_headers)
    assert accepted.status_code in (200, 201), accepted.text
    assert "\\u0000" not in accepted.text and "\x00" not in accepted.text


def test_the_conflict_response_is_documented_in_the_api_schema(client):
    schema = client.get("/openapi.json").json()
    patch = schema["paths"]["/api/v1/cv-documents/{document_id}"]["patch"]
    assert "409" in patch["responses"]


def test_a_tall_four_piece_entry_prints_its_middle_bullet_once(client, auth_headers):
    bullets = [
        f"Bullet {name}: " + " ".join(f"delivered outcome {n} for a large customer programme." for n in range(18))
        for name in ("alpha", "bravo", "charlie")
    ]
    sections = [
        _section(
            "experience",
            "Experience",
            0,
            [_entry("tall", 0, heading="Staff Engineer", subheading="Acme", start_date="2019",
                    end_date="2024", bullets=bullets)],
        )
    ]
    document = _create(client, auth_headers, sections)

    flat = " ".join(_pdf_text(_pdf(client, auth_headers, document["id"])).split())

    for name in ("alpha", "bravo", "charlie"):
        assert flat.count(f"Bullet {name}:") == 1, name
