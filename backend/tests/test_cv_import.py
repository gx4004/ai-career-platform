import io

import fitz
import pytest
from docx import Document

from app.models.cv_document import CvDocument
from app.models.evidence_item import EvidenceItem
from app.schemas.cv_documents import CvImportAccept
from app.services.cv_documents import accept_import

PREFIX = "/api/v1/cv-documents/import"


def _text_resume() -> bytes:
    return (
        b"Summary\nPlatform engineer focused on reliable systems.\n\n"
        b"Experience\nSenior Engineer at Example Corp\n"
        b"Improved deployment reliability by 20%.\n\n"
        b"Skills\nPython\nPostgreSQL\n"
    )


def test_text_import_returns_reviewable_structured_proposal_without_persisting(
    client, auth_headers, db
):
    response = client.post(
        f"{PREFIX}/proposals",
        files={"file": ("resume.txt", io.BytesIO(_text_resume()), "text/plain")},
        headers=auth_headers,
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"] == "resume.txt"
    assert payload["sections"][0]["kind"] == "summary"
    assert payload["sections"][1]["kind"] == "experience"
    assert payload["sections"][1]["entries"][1]["claim"]["kind"] == "achievement"
    assert payload["sections"][1]["entries"][1]["claim"]["provenance"] == "imported"
    assert db.query(CvDocument).count() == 0
    assert db.query(EvidenceItem).count() == 0


def test_accept_reviewed_proposal_atomically_creates_document_and_unconfirmed_claims(
    client, auth_headers, db
):
    proposal = client.post(
        f"{PREFIX}/proposals",
        files={"file": ("resume.txt", io.BytesIO(_text_resume()), "text/plain")},
        headers=auth_headers,
    ).json()
    proposal["name"] = "Reviewed CV"
    proposal["sections"][1]["entries"][1]["body"] = "Improved safe releases by 20%."
    proposal["sections"][1]["entries"][1]["claim"]["content"] = {
        "statement": "Improved safe releases by 20%."
    }

    response = client.post(f"{PREFIX}/accept", json=proposal, headers=auth_headers)

    assert response.status_code == 201
    document = response.json()
    assert document["name"] == "Reviewed CV"
    evidence_id = document["sections"][1]["entries"][1]["evidence_item_id"]
    item = db.query(EvidenceItem).filter(EvidenceItem.id == evidence_id).one()
    assert item.provenance == "imported"
    assert item.confirmation_state == "unconfirmed"
    assert item.content == {"statement": "Improved safe releases by 20%."}

    replay = client.post(f"{PREFIX}/accept", json=proposal, headers=auth_headers)
    assert replay.status_code == 201
    assert replay.json()["id"] == document["id"]
    assert db.query(CvDocument).count() == 1
    assert db.query(EvidenceItem).count() == sum(
        entry["claim"] is not None
        for section in proposal["sections"]
        for entry in section["entries"]
    )


def test_discarding_proposal_requires_no_server_call_and_leaves_no_content(
    client, auth_headers, db
):
    response = client.post(
        f"{PREFIX}/proposals",
        files={"file": ("resume.txt", io.BytesIO(_text_resume()), "text/plain")},
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert db.query(CvDocument).count() == 0
    assert db.query(EvidenceItem).count() == 0


def test_accept_rolls_back_document_and_claims_together(db, test_user, monkeypatch):
    body = CvImportAccept.model_validate(
        {
            "filename": "resume.txt",
            "import_id": "123e4567-e89b-42d3-a456-426614174000",
            "name": "Atomic CV",
            "warnings": [],
            "sections": [
                {
                    "id": "summary",
                    "kind": "summary",
                    "title": "Summary",
                    "visible": True,
                    "position": 0,
                    "entries": [
                        {
                            "id": "entry",
                            "body": "Synthetic claim",
                            "position": 0,
                            "claim": {
                                "kind": "experience",
                                "content": {"statement": "Synthetic claim"},
                                "provenance": "imported",
                            },
                        }
                    ],
                }
            ],
        }
    )

    def fail_commit():
        raise RuntimeError("synthetic database failure")

    monkeypatch.setattr(db, "commit", fail_commit)
    with pytest.raises(RuntimeError, match="synthetic database failure"):
        accept_import(db, test_user.id, body)
    assert db.query(CvDocument).count() == 0
    assert db.query(EvidenceItem).count() == 0


def test_import_contract_requires_fields_that_the_zod_mirror_requires():
    with pytest.raises(ValueError):
        CvImportAccept.model_validate(
            {
                "filename": "resume.txt",
                "import_id": "123e4567-e89b-42d3-a456-426614174000",
                "name": "Incomplete",
                "sections": [],
            }
        )
    with pytest.raises(ValueError):
        CvImportAccept.model_validate(
            {
                "filename": "resume.txt",
                "import_id": "123e4567e89b42d3a456426614174000",
                "name": "Noncanonical UUID",
                "warnings": [],
                "sections": [],
            }
        )


def test_import_validation_errors_are_actionable_without_echoing_content(
    client, auth_headers, monkeypatch
):
    from app.services import cv_upload

    monkeypatch.setattr(cv_upload, "MAX_CV_SIZE", 8)
    secret = b"private-resume-content"
    response = client.post(
        f"{PREFIX}/proposals",
        files={"file": ("resume.txt", io.BytesIO(secret), "text/plain")},
        headers=auth_headers,
    )

    assert response.status_code == 413
    assert response.json()["detail"] == "The uploaded file is larger than the 10 MB limit."
    assert secret.decode() not in response.text


def test_pdf_and_docx_imports_use_the_same_structured_contract(client, auth_headers):
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Skills\nPython")
    pdf_bytes = pdf.tobytes()
    pdf.close()

    docx_buffer = io.BytesIO()
    docx = Document()
    docx.add_paragraph("Education")
    table = docx.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Synthetic University"
    docx.add_paragraph("Skills")
    docx.add_paragraph("Python")
    docx.save(docx_buffer)

    cases = [
        ("resume.pdf", pdf_bytes, "application/pdf", "skills"),
        (
            "resume.docx",
            docx_buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "education",
        ),
    ]
    for filename, content, mime, expected_kind in cases:
        response = client.post(
            f"{PREFIX}/proposals",
            files={"file": (filename, io.BytesIO(content), mime)},
            headers=auth_headers,
        )
        assert response.status_code == 200
        assert response.json()["sections"][0]["kind"] == expected_kind
        assert response.json()["sections"][0]["entries"][0]["claim"]["kind"] in {
            "skill",
            "education",
        }
        if filename.endswith(".docx"):
            assert [section["kind"] for section in response.json()["sections"]] == [
                "education",
                "skills",
            ]
            assert response.json()["sections"][0]["entries"][0]["body"] == (
                "Synthetic University"
            )


def test_docx_zip_bomb_limit_returns_actionable_error(client, auth_headers, monkeypatch):
    from app.services import cv_upload

    output = io.BytesIO()
    document = Document()
    document.add_paragraph("Skills")
    document.add_paragraph("A" * 10_000)
    document.save(output)
    monkeypatch.setattr(cv_upload, "MAX_DOCX_COMPRESSION_RATIO", 1)

    response = client.post(
        f"{PREFIX}/proposals",
        files={
            "file": (
                "resume.docx",
                io.BytesIO(output.getvalue()),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=auth_headers,
    )
    assert response.status_code == 413
    assert response.json()["detail"] == (
        "The DOCX expands beyond safe processing limits. Try a simpler document."
    )
    assert "AAAA" not in response.text


def test_parser_timeout_returns_safe_actionable_error(client, auth_headers, monkeypatch):
    from app.services.cv_parser_process import CvParserProcessRejected

    async def time_out(*args, **kwargs):
        raise CvParserProcessRejected("Parser timed out")

    monkeypatch.setattr("app.routers.cv_documents.parse_cv_import_isolated", time_out)
    secret = b"Summary\nprivate timeout content"
    response = client.post(
        f"{PREFIX}/proposals",
        files={"file": ("resume.txt", io.BytesIO(secret), "text/plain")},
        headers=auth_headers,
    )
    assert response.status_code == 400
    assert response.json()["detail"] == "Import took too long. Try a smaller or simpler document."
    assert secret.decode() not in response.text


def test_encrypted_pdf_returns_safe_actionable_error(client, auth_headers):
    document = fitz.open()
    document.new_page().insert_text((72, 72), "private encrypted content")
    encrypted = document.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner",
        user_pw="secret",
    )
    document.close()

    response = client.post(
        f"{PREFIX}/proposals",
        files={"file": ("resume.pdf", io.BytesIO(encrypted), "application/pdf")},
        headers=auth_headers,
    )
    assert response.status_code == 400
    assert "unencrypted PDF" in response.json()["detail"]
    assert "private encrypted content" not in response.text
