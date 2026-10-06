"""CV header round-trip, seeding a CV from Evidence, and oversize-evidence
robustness (B5)."""

import io

import fitz
from docx import Document

PREFIX = "/api/v1/cv-documents"
EVIDENCE = "/api/v1/evidence-profile/items"


def _fact(client, auth_headers, kind, content):
    response = client.post(
        EVIDENCE,
        json={"kind": kind, "content": content, "provenance": "user-entered"},
        headers=auth_headers,
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _pdf_text(client, auth_headers, document_id):
    response = client.get(f"{PREFIX}/{document_id}/artifacts/pdf", headers=auth_headers)
    assert response.status_code == 200
    return "\n".join(page.get_text() for page in fitz.open(stream=response.content, filetype="pdf"))


def _docx_text(client, auth_headers, document_id):
    response = client.get(f"{PREFIX}/{document_id}/artifacts/docx", headers=auth_headers)
    assert response.status_code == 200
    return "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)


HEADER = {
    "name": "Casey Morgan",
    "headline": "Senior Backend Engineer",
    "email": "casey.morgan@example.com",
    "phone": "+49 30 5550134",
    "location": "Berlin, Germany",
    "links": ["https://github.com/caseymorgan"],
}


def test_header_round_trips_through_the_api_and_both_exports(client, auth_headers):
    created = client.post(
        PREFIX, json={"name": "internal-name-xyz", "header": HEADER, "sections": []},
        headers=auth_headers,
    )
    assert created.status_code == 201, created.text
    document_id = created.json()["id"]
    assert created.json()["header"] == HEADER
    assert client.get(f"{PREFIX}/{document_id}", headers=auth_headers).json()["header"] == HEADER
    assert client.get(PREFIX, headers=auth_headers).json()["items"][0]["header"] == HEADER

    for text in (
        _pdf_text(client, auth_headers, document_id),
        _docx_text(client, auth_headers, document_id),
    ):
        assert "Casey Morgan" in text
        assert "Senior Backend Engineer" in text
        assert "casey.morgan@example.com" in text
        assert "Berlin, Germany" in text
        assert "internal-name-xyz" not in text


def test_header_can_be_edited_and_cleared_with_patch(client, auth_headers):
    document_id = client.post(
        PREFIX, json={"name": "cv", "header": HEADER, "sections": []}, headers=auth_headers
    ).json()["id"]

    patched = client.patch(
        f"{PREFIX}/{document_id}",
        json={"header": {"name": "Casey M. Morgan", "headline": None}},
        headers=auth_headers,
    )

    assert patched.status_code == 200, patched.text
    header = patched.json()["header"]
    assert header["name"] == "Casey M. Morgan"
    assert header["headline"] is None and header["email"] is None and header["links"] == []
    assert "Casey M. Morgan" in _pdf_text(client, auth_headers, document_id)


def test_new_cv_header_defaults_to_the_account_name_not_the_document_name(
    client, auth_headers
):
    created = client.post(
        PREFIX, json={"name": "My CV", "sections": []}, headers=auth_headers
    ).json()

    assert created["header"]["name"] == "Test User"
    text = _pdf_text(client, auth_headers, created["id"])
    assert "Test User" in text
    assert "My CV" not in text


def test_header_rejects_oversize_values_with_422(client, auth_headers):
    response = client.post(
        PREFIX,
        json={"name": "cv", "header": {"name": "x" * 500}, "sections": []},
        headers=auth_headers,
    )
    assert response.status_code == 422


def test_seed_groups_facts_into_one_human_section_per_kind(client, auth_headers):
    skills = [
        _fact(client, auth_headers, "skill", {"statement": name})
        for name in ("Python", "PostgreSQL", "Kubernetes")
    ]
    role = _fact(
        client,
        auth_headers,
        "experience",
        {
            "title": "Senior Backend Engineer",
            "company": "Northwind Labs",
            "start_date": "Mar 2021",
            "end_date": "Present",
            "summary": "Leads the platform team.",
        },
    )
    achievement = _fact(
        client, auth_headers, "achievement", {"text": "Cut p95 latency by 38% on the ingest API."}
    )
    role2 = _fact(
        client,
        auth_headers,
        "experience",
        {"role": "Backend Engineer", "employer": "Acme Corp", "statement": "Built billing."},
    )

    response = client.post(
        PREFIX,
        json={
            "name": "From evidence",
            "sections": [],
            "seed_evidence_item_ids": [skills[0], role, skills[1], achievement, role2, skills[2]],
        },
        headers=auth_headers,
    )

    assert response.status_code == 201, response.text
    sections = response.json()["sections"]
    assert [s["title"] for s in sections] == ["Experience", "Achievements", "Skills"]
    assert [s["kind"] for s in sections] == ["experience", "achievements", "skills"]
    experience, achievements, skills_section = sections
    assert [e["heading"] for e in experience["entries"]] == [
        "Senior Backend Engineer",
        "Backend Engineer",
    ]
    first = experience["entries"][0]
    assert first["subheading"] == "Northwind Labs"
    assert (first["start_date"], first["end_date"]) == ("Mar 2021", "Present")
    assert first["body"] == "Leads the platform team."
    assert experience["entries"][1]["body"] == "Built billing."
    assert achievements["entries"][0]["body"] == "Cut p95 latency by 38% on the ingest API."
    # B16: skills read as one comma-joined line, in the order they were picked.
    assert [e["body"] for e in skills_section["entries"]] == ["Python, PostgreSQL, Kubernetes"]
    assert [e["position"] for e in skills_section["entries"]] == [0]
    assert experience["entries"][0]["evidence_item_id"] == role


def test_imported_role_evidence_seeds_a_role_with_its_bullets(client, auth_headers):
    role = _fact(
        client,
        auth_headers,
        "experience",
        {
            "title": "Senior Backend Engineer",
            "company": "Northwind Labs",
            "start_date": "Mar 2021",
            "highlights": "Cut p95 latency by 38%.\nLed 4 engineers.",
        },
    )

    created = client.post(
        PREFIX, json={"name": "seeded", "sections": [], "seed_evidence_item_ids": [role]},
        headers=auth_headers,
    ).json()

    entry = created["sections"][0]["entries"][0]
    assert entry["bullets"] == ["Cut p95 latency by 38%.", "Led 4 engineers."]
    assert entry["heading"] == "Senior Backend Engineer"
    assert entry["start_date"] == "Mar 2021" and entry["end_date"] is None


def test_oversize_bullets_are_rejected_before_they_can_break_the_list(client, auth_headers):
    section = {
        "id": "s",
        "kind": "experience",
        "title": "Experience",
        "position": 0,
        "entries": [
            {
                "id": "e",
                "evidence_item_id": None,
                "body": "x",
                "position": 0,
                "heading": "Role",
                "bullets": ["w" * 900 for _ in range(10)],
            }
        ],
    }

    created = client.post(
        PREFIX, json={"name": "cv", "sections": [section]}, headers=auth_headers
    )

    assert created.status_code == 422
    assert client.get(PREFIX, headers=auth_headers).status_code == 200


def test_seeded_cv_never_contains_a_dict_repr_in_the_paper_or_exports(client, auth_headers):
    ids = [
        _fact(client, auth_headers, "achievement", {"text": "Cut p95 latency by 38%."}),
        _fact(
            client,
            auth_headers,
            "experience",
            {"title": "Senior Backend Engineer", "company": "Northwind Labs"},
        ),
        _fact(client, auth_headers, "skill", {"name": "Python", "level": "expert"}),
    ]

    created = client.post(
        PREFIX, json={"name": "seeded", "sections": [], "seed_evidence_item_ids": ids},
        headers=auth_headers,
    ).json()

    bodies = [e["body"] for s in created["sections"] for e in s["entries"]]
    for body in bodies:
        assert "{" not in body and "'" not in body.split(" ")[0]
    assert "Cut p95 latency by 38%." in bodies
    assert "Python" in bodies
    pdf = _pdf_text(client, auth_headers, created["id"])
    docx = _docx_text(client, auth_headers, created["id"])
    for text in (pdf, docx):
        assert "{" not in text and "}" not in text
        assert "Senior Backend Engineer" in text
    # An experience with no description does not repeat its heading as body text.
    assert pdf.count("Senior Backend Engineer") == 1


def test_oversize_evidence_cannot_break_creation_or_listing(
    client, auth_headers, db, test_user
):
    # The evidence API now bounds a fact, so an oversize one can only be a row stored
    # before that bound; seed it directly.
    from app.models.evidence_item import EvidenceItem

    legacy = EvidenceItem(
        user_id=test_user.id,
        kind="achievement",
        content={"statement": "y" * 6000},
        provenance="user-entered",
        confirmation_state="confirmed",
    )
    db.add(legacy)
    db.commit()
    big = legacy.id
    small = _fact(client, auth_headers, "skill", {"statement": "Python"})

    created = client.post(
        PREFIX,
        json={"name": "big", "sections": [], "seed_evidence_item_ids": [big, small]},
        headers=auth_headers,
    )

    assert created.status_code == 201, created.text
    entries = created.json()["sections"][0]["entries"]
    assert 0 < len(entries[0]["body"]) <= 5000
    listing = client.get(PREFIX, headers=auth_headers)
    assert listing.status_code == 200
    assert [item["name"] for item in listing.json()["items"]] == ["big"]
    assert client.get(f"{PREFIX}/{created.json()['id']}/artifacts/pdf", headers=auth_headers).status_code == 200


def test_listing_survives_a_stored_document_that_violates_the_current_limits(
    client, auth_headers, db, test_user
):
    from app.models.cv_document import CvDocument

    good = client.post(
        PREFIX, json={"name": "good", "sections": []}, headers=auth_headers
    ).json()
    db.add(
        CvDocument(
            user_id=test_user.id,
            name="legacy-bad",
            sections=[
                {
                    "id": "s",
                    "kind": "achievements",
                    "title": "Achievements",
                    "visible": True,
                    "position": 0,
                    "entries": [
                        {
                            "id": "e",
                            "evidence_item_id": None,
                            "body": "z" * 9000,
                            "position": 0,
                        }
                    ],
                }
            ],
        )
    )
    db.commit()

    listing = client.get(PREFIX, headers=auth_headers)

    assert listing.status_code == 200
    names = {item["name"] for item in listing.json()["items"]}
    assert names == {"good", "legacy-bad"}
    assert good["id"] in {item["id"] for item in listing.json()["items"]}
    bad = next(item for item in listing.json()["items"] if item["name"] == "legacy-bad")
    assert len(bad["sections"][0]["entries"][0]["body"]) <= 5000
