"""DOCX per template and the plain-text export (spec D8, T10)."""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from docx import Document
from docx.shared import Mm

from app.schemas.cv_documents import CvStyle
from app.services.cv_docx import TWO_COLUMN_DOCX_NOTE
from app.services.cv_html import load_manifest
from app.services.cv_rendering import (
    TEMPLATES,
    build_render_model,
    render_docx,
    render_txt,
    style_catalog,
)
from tests.cv_fixtures import ALL

PREFIX = "/api/v1/cv-documents"


def _model(fixture: str, template_id: str, **style):
    return build_render_model(ALL[fixture](), template_id, CvStyle(template_id=template_id, **style))


def _part(docx: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        return archive.read(name).decode()


@pytest.mark.parametrize("fixture", list(ALL))
@pytest.mark.parametrize("template_id", list(TEMPLATES))
def test_docx_for_every_template_and_fixture(template_id, fixture):
    model = _model(fixture, template_id)
    docx = render_docx(model)
    assert docx == render_docx(model)  # byte-stable
    manifest = load_manifest(template_id)
    parsed = Document(io.BytesIO(docx))
    styles_xml, document_xml = _part(docx, "word/styles.xml"), _part(docx, "word/document.xml")

    # The template's fonts by family name, accent on the headings, one column, no layout tables.
    for family in manifest.typefaces.values():
        assert f'w:ascii="{family}"' in styles_xml
        assert f'w:name="{family}"' in _part(docx, "word/fontTable.xml")
    accent = (manifest.default_accent or "#111827")[1:].upper()
    assert accent in styles_xml.upper() or accent in document_xml.upper()
    assert "<w:cols" not in document_xml.replace('<w:cols w:space="720"/>', "")
    assert not parsed.tables

    # Section order and every entry survive, headings are real Word headings.
    headings = [p.text for p in parsed.paragraphs if p.style.name == "Heading 1"]
    assert headings == [s.title for s in model.sections]
    text = "\n".join(p.text for p in parsed.paragraphs)
    for section in model.sections:
        for entry in section.entries:
            for piece in (entry.heading, entry.subheading, entry.paragraph, *entry.bullets):
                if piece:
                    assert piece in text
    assert model.header.title in text
    if manifest.title_align == "center":
        assert parsed.paragraphs[0].alignment == 1
    assert any(p.style.name == "List Bullet" for p in parsed.paragraphs)
    assert any(p.style.name == "Heading 2" for p in parsed.paragraphs)
    # Cyrillic and accents are kept as written (NFC).
    for char in model.header.title:
        assert char in text


def test_page_size_and_margins_follow_the_style():
    letter = Document(io.BytesIO(render_docx(_model("maya", "classic", page_size="letter"))))
    a4 = Document(io.BytesIO(render_docx(_model("maya", "classic"))))
    assert abs(letter.sections[0].page_width - Mm(215.9)) < Mm(0.5)
    assert abs(a4.sections[0].page_height - Mm(297)) < Mm(0.5)
    assert abs(a4.sections[0].left_margin - Mm(17)) < Mm(0.5)


def test_typeface_override_and_density_reach_the_docx():
    docx = render_docx(_model("maya", "classic", font_id="lora", density="compact"))
    assert 'w:ascii="Lora"' in _part(docx, "word/styles.xml")
    spacious = render_docx(_model("maya", "classic", density="spacious"))
    assert _part(docx, "word/styles.xml") != _part(spacious, "word/styles.xml")


def test_dates_use_a_right_tab_and_links_are_hyperlinks():
    docx = render_docx(_model("maya", "classic"))
    document_xml = _part(docx, "word/document.xml")
    assert 'w:val="right"' in document_xml and "<w:tab/>" in document_xml
    rels = _part(docx, "word/_rels/document.xml.rels")
    assert "mailto:maya.lindqvist@example.com" in rels
    assert 'Target="https://linkedin.com/in/maya-lindqvist-example"' in rels
    assert "<w:hyperlink" in document_xml


def test_two_column_templates_print_one_column_and_say_so():
    catalog = {t.id: t for t in style_catalog().templates}
    for template_id, template in TEMPLATES.items():
        note = catalog[template_id].docx_note
        assert note == (TWO_COLUMN_DOCX_NOTE if template.two_column else None)
    assert TEMPLATES["lagoon"].two_column and catalog["classic"].docx_note is None


@pytest.mark.parametrize("fixture", list(ALL))
@pytest.mark.parametrize("template_id", ["classic", "lagoon"])
def test_txt_follows_section_order_and_omits_empty_sections(template_id, fixture):
    model = _model(fixture, template_id)
    raw = render_txt(model)
    text = raw.decode("utf-8")
    assert text.endswith("\n") and "\r" not in text
    lines = text.splitlines()
    assert lines[0] == model.header.title
    positions = [lines.index(s.title.upper()) for s in model.sections]
    assert positions == sorted(positions)
    for position, section in zip(positions, model.sections, strict=True):
        assert lines[position + 1] == "-" * len(section.title.upper())
        for entry in section.entries:
            for bullet in entry.bullets:
                assert f"- {bullet}" in lines
    assert re.search(r"^Frontend Engineer, Tulip Pay \(Amsterdam\) \| Mar 2022 – Present$", text, re.M) or fixture != "maya"


def test_txt_skills_are_comma_separated_and_empty_sections_are_dropped():
    cv = ALL["maya"]()
    cv.sections.append(
        {"id": "z", "kind": "custom", "title": "Awards", "visible": True, "position": 9, "entries": []}
    )
    model = build_render_model(cv, "classic", CvStyle())
    text = render_txt(model).decode()
    assert "TypeScript, JavaScript, React, Next.js" in text
    assert "AWARDS" not in text


# -- endpoints -------------------------------------------------------------------------------


def test_txt_endpoint_headers_and_limit_is_separate(client, auth_headers):
    sections = [{"id": "s", "kind": "summary", "title": "Summary", "visible": True, "position": 0,
                 "entries": [{"id": "e", "evidence_item_id": None, "position": 0, "body": "Hello"}]}]
    created = client.post(PREFIX, json={"name": "Txt CV", "sections": sections}, headers=auth_headers)
    assert created.status_code in (200, 201), created.text
    doc_id = created.json()["id"]
    response = client.get(f"{PREFIX}/{doc_id}/artifacts/txt", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/plain")
    assert response.headers["content-disposition"].endswith('.txt"')
    assert "SUMMARY" in response.text and "Hello" in response.text
    assert client.get(f"{PREFIX}/{doc_id}/artifacts/txt").status_code in (401, 403)
