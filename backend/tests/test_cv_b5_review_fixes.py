"""Review fixes for B5: legacy oversize rows on every mutation, Word Heading 2 roles,
bounded URL scanning, header stoplist, three-line role headers, seeded summaries."""

import io
import logging
import time

import pytest
from docx import Document

from app.services.cv_parser import parse_cv_import

PREFIX = "/api/v1/cv-documents"
EVIDENCE = "/api/v1/evidence-profile/items"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _legacy_oversize_document(db, user, *, with_variant=True):
    from app.models.cv_document import CvDocument, CvVariant

    sections = [
        {
            "id": "s",
            "kind": "achievements",
            "title": "Achievements",
            "visible": True,
            "position": 0,
            "entries": [
                {"id": "e", "evidence_item_id": None, "body": "z" * 9000, "position": 0}
            ],
        }
    ]
    document = CvDocument(user_id=user.id, name="legacy-bad", sections=sections)
    if with_variant:
        document.variants.append(CvVariant(name="Base", sections=sections))
    db.add(document)
    db.commit()
    return document


def test_mutations_on_a_legacy_oversize_row_do_not_500(client, auth_headers, db, test_user):
    document = _legacy_oversize_document(db, test_user)
    url = f"{PREFIX}/{document.id}"

    renamed = client.patch(url, json={"name": "renamed"}, headers=auth_headers)
    assert renamed.status_code == 200, renamed.text
    assert len(renamed.json()["sections"][0]["entries"][0]["body"]) <= 5000

    styled = client.patch(
        url,
        json={"style": {"template_id": "classic", "font_id": "lato",
                        "accent_color": "#111827", "density": "normal", "ats_mode": False}},
        headers=auth_headers,
    )
    assert styled.status_code == 200, styled.text

    headed = client.patch(url, json={"header": {"name": "Casey"}}, headers=auth_headers)
    assert headed.status_code == 200, headed.text
    assert headed.json()["header"]["name"] == "Casey"

    variant = client.post(f"{url}/variants", json={"name": "Snap"}, headers=auth_headers)
    assert variant.status_code == 201, variant.text
    assert len(variant.json()["sections"][0]["entries"][0]["body"]) <= 5000

    base = next(v for v in client.get(url, headers=auth_headers).json()["variants"] if v["name"] == "Base")
    restored = client.post(f"{url}/variants/{base['id']}/restore", headers=auth_headers)
    assert restored.status_code == 200, restored.text


def test_unreadable_list_row_is_a_recoverable_stub_and_logs_no_content(
    client, auth_headers, db, test_user, caplog, monkeypatch
):
    document = _legacy_oversize_document(db, test_user)
    from app.services import cv_documents

    # Simulate a row the clamp cannot repair either.
    def broken(_sections):
        raise ValueError("secret cv content 9000 z")

    monkeypatch.setattr(cv_documents, "_clamp_sections", broken)
    with caplog.at_level(logging.DEBUG):
        listing = client.get(PREFIX, headers=auth_headers)

    assert listing.status_code == 200
    (item,) = listing.json()["items"]
    assert item["id"] == document.id and item["sections"] == []
    assert "secret cv content" not in caplog.text and "zzzz" not in caplog.text
    assert client.delete(f"{PREFIX}/{document.id}", headers=auth_headers).status_code == 204


def _docx(build):
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_heading_one_sections_with_heading_two_roles(client, auth_headers):
    def build(d):
        d.add_heading("Casey Morgan", 0)
        d.add_paragraph("casey.morgan@example.com | +49 30 5550134")
        d.add_heading("Experience", 1)
        d.add_heading("Senior Engineer, Acme Corp", 2)
        d.add_paragraph("Mar 2021 – Present")
        d.add_paragraph("Cut p95 latency by 38%.", style="List Bullet")
        d.add_paragraph("Led 4 engineers.", style="List Bullet")
        d.add_heading("Engineer, Globex", 2)
        d.add_paragraph("Jan 2018 – Feb 2021")
        d.add_paragraph("Built the billing service.", style="List Bullet")
        d.add_heading("Skills", 1)
        d.add_paragraph("Python")

    proposal = parse_cv_import(_docx(build), "cv.docx", "docx").model_dump()

    assert [s["kind"] for s in proposal["sections"]] == ["experience", "skills"]
    first, second = proposal["sections"][0]["entries"]
    assert (first["heading"], first["subheading"]) == ("Senior Engineer", "Acme Corp")
    assert (first["start_date"], first["end_date"]) == ("Mar 2021", "Present")
    assert first["bullets"] == ["Cut p95 latency by 38%.", "Led 4 engineers."]
    assert (second["heading"], second["subheading"]) == ("Engineer", "Globex")
    assert second["start_date"] == "Jan 2018"
    assert second["bullets"] == ["Built the billing service."]


def test_a_bare_date_line_never_becomes_a_role_heading():
    text = "Experience\nMar 2021 – Present\nBuilt things.\n"
    proposal = parse_cv_import(text.encode(), "cv.txt", "txt").model_dump()
    for entry in proposal["sections"][0]["entries"]:
        assert entry.get("heading") not in {"Mar", "Mar 2021"}


def test_long_dotted_tokens_import_quickly():
    body = "a." * 2500
    text = "Awards\n" + "\n".join([body] * 100) + "\n"
    started = time.monotonic()
    proposal = parse_cv_import(text.encode(), "cv.txt", "txt")
    assert time.monotonic() - started < 3
    assert proposal.sections


@pytest.mark.parametrize("banner", ["CURRICULUM VITAE", "Resume", "Curriculum Vitae", "CV"])
def test_document_title_banner_is_not_the_name(banner):
    text = f"{banner}\nCasey Morgan\ncasey.morgan@example.com\n\nSkills\nPython\n"
    header = parse_cv_import(text.encode(), "cv.txt", "txt").header
    assert header.name == "Casey Morgan"
    assert header.headline in (None, "")


def test_three_line_role_header_is_one_role():
    text = (
        "Experience\nSenior Engineer\nStripe\nJan 2020 – Dec 2023\n"
        "• Cut costs by 20%.\n\nEducation\nStanford University\nBS Computer Science\n2012 – 2016\n"
    )
    proposal = parse_cv_import(text.encode(), "cv.txt", "txt").model_dump()
    (role,) = proposal["sections"][0]["entries"]
    assert (role["heading"], role["subheading"]) == ("Senior Engineer", "Stripe")
    assert (role["start_date"], role["end_date"]) == ("Jan 2020", "Dec 2023")
    assert role["bullets"] == ["Cut costs by 20%."]
    (edu,) = proposal["sections"][1]["entries"]
    assert edu["start_date"] == "2012" and edu["end_date"] == "2016"
    assert edu["heading"] and edu["subheading"]


def test_seeded_experience_keeps_its_summary_next_to_its_bullets(client, auth_headers):
    created_fact = client.post(
        EVIDENCE,
        json={
            "kind": "experience",
            "content": {
                "title": "Engineer",
                "company": "Acme",
                "summary": "Owns the platform.",
                "highlights": "Shipped X\nCut cost 20%",
            },
            "provenance": "user-entered",
        },
        headers=auth_headers,
    )
    assert created_fact.status_code == 201, created_fact.text
    created = client.post(
        PREFIX,
        json={"name": "s", "sections": [], "seed_evidence_item_ids": [created_fact.json()["id"]]},
        headers=auth_headers,
    )
    assert created.status_code == 201, created.text
    entry = created.json()["sections"][0]["entries"][0]
    assert "Owns the platform." in "\n".join(entry["bullets"])
    assert "Shipped X" in entry["bullets"]
