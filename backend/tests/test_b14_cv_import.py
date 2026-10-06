"""B14 item 10: CV import edge cases, through POST /cv-documents/import/proposals."""

from __future__ import annotations

import io

import pytest
from docx import Document

PREFIX = "/api/v1/cv-documents"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _propose(client, headers, filename, content, mime):
    response = client.post(
        f"{PREFIX}/import/proposals",
        files={"file": (filename, io.BytesIO(content), mime)},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return response.json()


def _text(client, headers, text: str):
    return _propose(client, headers, "cv.txt", text.encode(), "text/plain")


def _by_kind(proposal):
    return {section["kind"]: section for section in proposal["sections"]}


def test_a_us_phone_number_keeps_its_parentheses(client, auth_headers):
    proposal = _text(
        client,
        auth_headers,
        "Jordan Rivera\njordan.rivera@example.com | (555) 123-4567 | Austin, TX\n\nSkills\nPython\n",
    )

    assert proposal["header"]["phone"] == "(555) 123-4567"
    assert proposal["header"]["location"] == "Austin, TX"


def test_a_non_place_segment_on_the_contact_line_is_not_the_location(client, auth_headers):
    proposal = _text(
        client,
        auth_headers,
        "Jordan Rivera\njordan.rivera@example.com | (555) 123-4567 | Senior Backend Engineer\n\nSkills\nPython\n",
    )

    header = proposal["header"]
    assert header["location"] is None
    assert header["phone"] == "(555) 123-4567"
    assert header["headline"] == "Senior Backend Engineer"


def test_an_inline_skills_line_is_its_own_section_not_part_of_the_last_bullet(client, auth_headers):
    proposal = _text(
        client,
        auth_headers,
        "Jordan Rivera\n\nExperience\nBackend Engineer, Acme Corp (2019 - 2024)\n"
        "- Built the billing service.\n\nTechnical Skills: Go, Rust, SQL\n",
    )

    role = _by_kind(proposal)["experience"]["entries"][0]
    assert role["bullets"] == ["Built the billing service."]
    skills = _by_kind(proposal)["skills"]
    assert skills["title"] == "Technical Skills"
    assert [entry["body"] for entry in skills["entries"]] == ["Go, Rust, SQL"]


def test_a_wrapped_bullet_does_not_continue_across_a_blank_line(client, auth_headers):
    proposal = _text(
        client,
        auth_headers,
        "Jordan Rivera\n\nExperience\nBackend Engineer, Acme Corp (2019 - 2024)\n"
        "- Built the billing service for\n  three product lines.\n\nOn-call lead for the payments team\n",
    )

    role = _by_kind(proposal)["experience"]["entries"][0]
    assert role["bullets"] == [
        "Built the billing service for three product lines.",
        "On-call lead for the payments team",
    ]


def test_docx_heading_role_with_a_company_and_dates_line_is_one_entry(client, auth_headers):
    document = Document()
    document.add_heading("Jordan Rivera", 0)
    document.add_heading("Experience", 1)
    document.add_heading("Senior Backend Engineer", 2)
    document.add_paragraph("Northwind Labs | Mar 2021 – Present")
    document.add_paragraph("Cut p95 latency by 38%.", style="List Bullet")
    document.add_heading("Backend Engineer", 2)
    document.add_paragraph("Acme Corp | Jan 2018 – Feb 2021")
    document.add_paragraph("Built the billing service.", style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)

    proposal = _propose(client, auth_headers, "cv.docx", buffer.getvalue(), DOCX_MIME)

    roles = _by_kind(proposal)["experience"]["entries"]
    assert [(r["heading"], r.get("subheading"), r.get("start_date"), r.get("end_date")) for r in roles] == [
        ("Senior Backend Engineer", "Northwind Labs", "Mar 2021", "Present"),
        ("Backend Engineer", "Acme Corp", "Jan 2018", "Feb 2021"),
    ]
    assert roles[0]["bullets"] == ["Cut p95 latency by 38%."]


@pytest.mark.parametrize(
    "place",
    [
        "San Francisco Bay Area",
        "New York City",
        "New York, NY 10001",
        "Rio de Janeiro",
        "Remote (EU)",
        "Remote (US)",
    ],
)
def test_a_place_on_the_contact_line_stays_the_location(client, auth_headers, place):
    proposal = _text(
        client, auth_headers, f"Jordan Rivera\njordan@example.com | {place}\n\nSkills\nPython\n"
    )

    assert proposal["header"]["location"] == place
    assert proposal["header"]["headline"] is None


def test_category_lines_inside_a_skills_section_stay_skills(client, auth_headers):
    proposal = _text(
        client,
        auth_headers,
        "Jordan Rivera\n\nSkills\nLanguages: Go, Python\nTools: Docker, Terraform\n"
        "Certifications: AWS Solutions Architect\n",
    )

    sections = proposal["sections"]
    assert [(section["kind"], section["title"]) for section in sections] == [("skills", "Skills")]
    entries = sections[0]["entries"]
    assert [entry["body"] for entry in entries] == [
        "Languages: Go, Python",
        "Tools: Docker, Terraform",
        "Certifications: AWS Solutions Architect",
    ]
    assert all(entry["claim"] and entry["claim"]["kind"] == "skill" for entry in entries)

