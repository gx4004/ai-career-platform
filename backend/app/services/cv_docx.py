"""The CV as a Word document (docx) and as plain text, both from the render model.

DOCX design (docs/cv-templates-spec.md D8), within what Word does without layout tables:

* Fonts are the template's pair (or the typeface override) set by family name in the styles.
  Word does not ship them, so ``fontTable.xml`` names a generic family and an ``altName``
  (Georgia for serif, Arial for sans) that Word substitutes when the font is not installed.
  Fonts are NOT embedded: python-docx has no support for it, Word's embedding needs obfuscated
  ``.odttf`` parts that cannot be verified here without Word, and it would add 1 to 2 MB per file.
* One column for every template. Sidebar and rail templates keep their section order and get a
  coloured name band and tinted section headings instead (``docx_note`` tells the person).
* Real styles: Title (name), Subtitle (headline), Heading 1 (sections), Heading 2 (entries),
  List Bullet (bullets), so Word's navigation pane and ATS parsers see the structure.
* Dates are right-aligned with a tab stop; links are hyperlinks; no tables.

Bytes are stable: ZIP member order and timestamps are canonicalised.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import UTC, datetime

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

from app.schemas.cv_documents import CvRenderEntry, CvRenderModel
from app.services.cv_fonts import TYPEFACES_BY_NAME
from app.services.cv_html import (
    GLOBAL_DEFAULT_ACCENT,
    URL_RE,
    html_template_id,
    load_manifest,
)
from app.services.cv_style_tokens import INK, accent_tokens, contrast_ratio

TWO_COLUMN_DOCX_NOTE = "Word version uses a single column."

PAGE_MM = {"a4": (210.0, 297.0), "letter": (215.9, 279.4)}
# Two-column templates have a narrow outer margin for their sidebar; one column needs more.
MIN_ONE_COLUMN_MARGIN_MM = 16
GREY = "595959"
LINK_BLUE = "0563C1"
# Heading treatment of the single-column templates; two-column ones use the tinted band.
_CAPS = {"frame"}
_SMALL_CAPS = {"scholar"}

_HOSTLIKE = re.compile(r"^(https?://)?[\w-]+(\.[\w-]+)+(/\S*)?$", re.IGNORECASE)
_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def docx_note_for(two_column: bool) -> str | None:
    return TWO_COLUMN_DOCX_NOTE if two_column else None


def _generic(name: str) -> tuple[str, str]:
    """(w:family, altName) for a font family name: a serif or sans fallback Word has everywhere."""
    face = TYPEFACES_BY_NAME.get(name)
    if face is not None and face.category == "serif":
        return "roman", "Georgia"
    return "swiss", "Arial"


def _font_table(names: list[str]) -> bytes:
    fonts = []
    for name in dict.fromkeys(names):
        family, alt = _generic(name)
        fonts.append(
            f'<w:font w:name="{name}"><w:altName w:val="{alt}"/><w:family w:val="{family}"/>'
            '<w:pitch w:val="variable"/></w:font>'
        )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:fonts {nsdecls("w")}>{"".join(fonts)}</w:fonts>'
    ).encode()


def _set_fonts(rpr, name: str) -> None:
    """Name the font for every script (Cyrillic uses hAnsi; complex scripts use cs) and drop theme fonts."""
    fonts = rpr.get_or_add_rFonts()
    for attribute in list(fonts.attrib):
        del fonts.attrib[attribute]
    for attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attribute}"), name)


def _readable(color: str, background: str) -> str:
    return color if contrast_ratio(color, background) >= 4.5 else INK


def _border(element_name: str, color: str, size: int, space: int = 1):
    border = OxmlElement(element_name)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color.lstrip("#"))
    return border


def _shade(ppr, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill.lstrip("#"))
    ppr.append(shading)


def _set_borders(style_or_paragraph_ppr, borders: list) -> None:
    ppr = style_or_paragraph_ppr
    for existing in ppr.findall(qn("w:pBdr")):
        ppr.remove(existing)
    container = OxmlElement("w:pBdr")
    for border in borders:
        container.append(border)
    # Schema order: pBdr comes before shd, tabs and spacing.
    anchor = next(
        (ppr.find(qn(tag)) for tag in ("w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc") if ppr.find(qn(tag)) is not None),
        None,
    )
    if anchor is None:
        ppr.append(container)
    else:
        anchor.addprevious(container)


def _style_paragraph(style, *, before: float, after: float, line: float | None = None, keep_next=False) -> None:
    fmt = style.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    if line:
        fmt.line_spacing = line
    fmt.keep_with_next = keep_next or None


def _style_run(style, font: str, size: float, color: str | None, *, bold=False, italic=False) -> None:
    _set_fonts(style.element.get_or_add_rPr(), font)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if color:
        style.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def render_docx(model: CvRenderModel) -> bytes:
    tokens = model.tokens
    layout_id = html_template_id(model.template_id)
    manifest = load_manifest(layout_id)
    body_font = str(tokens["font_docx"])
    heading_font = str(tokens.get("font_docx_heading") or body_font)
    accent = str(tokens.get("accent") or manifest.default_accent or GLOBAL_DEFAULT_ACCENT).upper()
    palette = accent_tokens(accent)
    body_pt = int(tokens["body_size_pt"])
    heading_pt = int(tokens["heading_size_pt"])
    gap = int(tokens["gap_scale_pct"]) / 100
    two_column = bool(tokens.get("two_column"))
    centred = tokens["title_align"] == "center"

    doc = Document()
    page = doc.sections[0]
    width_mm, height_mm = PAGE_MM.get(str(tokens.get("page_size", "a4")), PAGE_MM["a4"])
    page.page_width, page.page_height = Mm(width_mm), Mm(height_mm)
    margins = [
        manifest.margin_top_mm, manifest.margin_right_mm, manifest.margin_bottom_mm, manifest.margin_left_mm,
    ]
    if two_column:
        margins = [max(m, MIN_ONE_COLUMN_MARGIN_MM) for m in margins]
    page.top_margin, page.right_margin, page.bottom_margin, page.left_margin = (Mm(m) for m in margins)
    content_width = Mm(width_mm - margins[1] - margins[3])

    core = doc.core_properties
    core.title = model.document_name
    core.author = "Career Workbench"
    core.created = core.modified = datetime(2000, 1, 1, tzinfo=UTC)

    styles = doc.styles
    # Document defaults would otherwise keep the template's theme fonts.
    defaults = styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        default_fonts = defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:rFonts')}")
        if default_fonts is not None:
            for attribute in list(default_fonts.attrib):
                del default_fonts.attrib[attribute]
            for attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
                default_fonts.set(qn(f"w:{attribute}"), body_font)

    ink = "1F2937"
    _style_run(styles["Normal"], body_font, body_pt, ink)
    _style_paragraph(styles["Normal"], before=0, after=0, line=1.12)

    banded = two_column
    band_text = palette["on_accent"].lstrip("#")
    name_color = band_text if banded else accent.lstrip("#")
    name_style = styles["Title"]
    _style_run(name_style, heading_font, heading_pt + 10, name_color, bold=True)
    title_ppr = name_style.element.get_or_add_pPr()
    for tag in ("w:pBdr", "w:contextualSpacing"):
        for found in title_ppr.findall(qn(tag)):
            title_ppr.remove(found)
    _style_paragraph(name_style, before=0, after=2)
    name_style.font.size = Pt(heading_pt + 10)
    for kerning in name_style.element.get_or_add_rPr().findall(qn("w:spacing")):
        kerning.getparent().remove(kerning)
    sub_style = styles["Subtitle"]
    _style_run(sub_style, body_font, body_pt + 2, band_text if banded else ink, italic=False)
    _style_paragraph(sub_style, before=0, after=3)
    for tag in ("w:spacing",):
        for found in sub_style.element.get_or_add_rPr().findall(qn(tag)):
            found.getparent().remove(found)

    contact = styles.add_style("Contact", 1)
    contact.base_style = styles["Normal"]
    _style_run(contact, body_font, body_pt - 0.5 if body_pt > 8 else body_pt, band_text if banded else GREY)
    _style_paragraph(contact, before=0, after=2)

    section_style = styles["Heading 1"]
    heading_color = _readable(accent, palette["accent_tint"] if banded else "#FFFFFF").lstrip("#")
    _style_run(section_style, heading_font, heading_pt, heading_color, bold=True)
    _style_paragraph(section_style, before=int(tokens["section_gap_pt"]) * 1.6 + 6, after=4, keep_next=True)
    ppr = section_style.element.get_or_add_pPr()
    if banded:
        _shade(ppr, palette["accent_tint"])
        _set_borders(ppr, [_border("w:left", accent, 24, 4)])
    else:
        _set_borders(ppr, [_border("w:bottom", accent, 6, 2)])
    rpr = section_style.element.get_or_add_rPr()
    if layout_id in _CAPS or banded:
        caps = OxmlElement("w:caps")
        rpr.append(caps)
    if layout_id in _SMALL_CAPS:
        rpr.append(OxmlElement("w:smallCaps"))

    entry_style = styles["Heading 2"]
    _style_run(entry_style, body_font, body_pt + 0.5, ink, bold=True)
    _style_paragraph(entry_style, before=5 * gap + 2, after=1, keep_next=True)
    detail = styles.add_style("Entry Detail", 1)
    detail.base_style = styles["Normal"]
    _style_run(detail, body_font, body_pt, GREY, italic=True)
    _style_paragraph(detail, before=0, after=2 * gap, keep_next=True)

    bullet = styles["List Bullet"]
    _style_run(bullet, body_font, body_pt, ink)
    _style_paragraph(bullet, before=0, after=2 * gap)
    bullet.paragraph_format.left_indent = Pt(16)
    bullet.paragraph_format.first_line_indent = Pt(-11)

    para = styles.add_style("Entry Text", 1)
    para.base_style = styles["Normal"]
    _style_paragraph(para, before=2 * gap, after=3 * gap)

    # Header block. Two-column templates get a coloured band (adjacent shaded paragraphs with the
    # same borders merge into one block in Word and LibreOffice).
    header_paragraphs = []

    def header_paragraph(style: str):
        p = doc.add_paragraph(style=style)
        p.alignment = 1 if centred else 0
        header_paragraphs.append(p)
        return p

    header_paragraph("Title").add_run(model.header.title)
    if model.header.headline:
        header_paragraph("Subtitle").add_run(model.header.headline)
    link_color = band_text if banded else LINK_BLUE
    if model.header.contact:
        line = header_paragraph("Contact")
        for index, item in enumerate(model.header.contact):
            if index:
                line.add_run("  |  ")
            _add_contact_item(line, item, link_color, underline=not banded)
    if banded:
        for p in header_paragraphs:
            p_ppr = p._p.get_or_add_pPr()
            _set_borders(
                p_ppr,
                [_border(side, accent, 4, 8) for side in ("w:top", "w:left", "w:bottom", "w:right")],
            )
            _shade(p_ppr, accent)
        # Space after the band comes from the first heading's space-before.
    for rendered in model.sections:
        heading = doc.add_paragraph(style="Heading 1")
        heading.add_run(rendered.title)
        for entry in rendered.entries:
            _add_entry(doc, entry, content_width)

    fonts = [body_font, heading_font]
    for part in doc.part.package.iter_parts():
        if str(part.partname) == "/word/fontTable.xml":
            part._blob = _font_table(fonts)

    raw = io.BytesIO()
    doc.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    final = io.BytesIO()
    with zipfile.ZipFile(final, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    source.close()
    return final.getvalue()


def _add_entry(doc, entry: CvRenderEntry, content_width) -> None:
    if entry.heading is not None:
        line = doc.add_paragraph(style="Heading 2")
        line.add_run(entry.heading)
        if entry.subheading:
            line.add_run(f", {entry.subheading}").bold = False
        if entry.dates:
            line.paragraph_format.tab_stops.add_tab_stop(content_width, WD_TAB_ALIGNMENT.RIGHT)
            line.add_run(f"\t{entry.dates}").bold = False
        if entry.location:
            doc.add_paragraph(entry.location, style="Entry Detail")
        for text in entry.bullets:
            _add_linked_text(doc.add_paragraph(style="List Bullet"), text)
    if entry.paragraph:
        p = doc.add_paragraph(style="Entry Text")
        p.paragraph_format.keep_together = True
        _add_linked_text(p, entry.paragraph)


def _add_contact_item(paragraph, item: str, color: str, *, underline: bool) -> None:
    if _EMAIL.match(item):
        _add_hyperlink(paragraph, item, f"mailto:{item}", color, underline)
    elif " " not in item and _HOSTLIKE.match(item):
        target = item if re.match(r"https?://", item, re.IGNORECASE) else f"https://{item}"
        _add_hyperlink(paragraph, item, target, color, underline)
    else:
        paragraph.add_run(item)


def _add_linked_text(paragraph, text: str) -> None:
    cursor = 0
    for match in URL_RE.finditer(text):
        paragraph.add_run(text[cursor : match.start()])
        _add_hyperlink(paragraph, match.group(), match.group(), LINK_BLUE, True)
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def _add_hyperlink(paragraph, text: str, target: str, color: str, underline: bool) -> None:
    relationship_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), color.lstrip("#"))
    properties.append(colour)
    if underline:
        line = OxmlElement("w:u")
        line.set(qn("w:val"), "single")
        properties.append(line)
    label = OxmlElement("w:t")
    label.text = text
    label.set(qn("xml:space"), "preserve")
    run.extend((properties, label))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


# ---------------------------------------------------------------------------- plain text


def render_txt(model: CvRenderModel) -> bytes:
    """The CV as UTF-8 plain text: header, then each section in order, empty sections omitted."""
    header = model.header
    lines: list[str] = [header.title]
    if header.headline:
        lines.append(header.headline)
    if header.contact:
        lines.append(" | ".join(header.contact))
    for section in model.sections:
        entries = [_txt_entry(section.kind, entry) for entry in section.entries]
        entries = [e for e in entries if e]
        if not entries:
            continue
        title = section.title.upper()
        lines += ["", title, "-" * len(title)]
        if section.kind == "skills":
            lines.append(", ".join(_skill_items(section.entries)))
        else:
            for index, block in enumerate(entries):
                if index:
                    lines.append("")
                lines += block
    return ("\n".join(lines) + "\n").encode("utf-8")


def _skill_items(entries: list[CvRenderEntry]) -> list[str]:
    items: list[str] = []
    for entry in entries:
        text = entry.heading or entry.paragraph or entry.text
        # Skills are saved as "A • B • C", "A, B" or one per entry; the export reads them as one list.
        for piece in re.split(r"\s*[•·|;\n]\s*", " ".join([text, *entry.bullets])):
            if piece.strip():
                items.append(piece.strip())
    return list(dict.fromkeys(items))


def _txt_entry(kind: str, entry: CvRenderEntry) -> list[str]:
    if entry.heading is None:
        return [entry.paragraph] if entry.paragraph else []
    title = entry.heading + (f", {entry.subheading}" if entry.subheading else "")
    if entry.location:
        title += f" ({entry.location})"
    if entry.dates:
        title += f" | {entry.dates}"
    lines = [title]
    if entry.paragraph:
        lines.append(entry.paragraph)
    lines += [f"- {bullet}" for bullet in entry.bullets]
    return lines
