import re
import uuid
from pathlib import Path

from app.schemas.cv_documents import CvImportProposal
from app.schemas.tools import ParsedCvResponse

MAX_PDF_PAGES = 100
MAX_EXTRACTED_CHARS = 2_000_000


class CvParserRejected(Exception):
    pass


def parse_cv(content: bytes, filename: str, ext: str) -> ParsedCvResponse:
    warnings: list[str] = []

    if ext == "pdf":
        text = _extract_pdf(content)
    elif ext == "docx":
        text = _extract_docx(content)
    else:
        raise ValueError(f"Unsupported format: {ext}")

    if not text.strip():
        warnings.append("No text could be extracted from this file")

    return ParsedCvResponse(
        filename=filename,
        extracted_text=text,
        chars_count=len(text),
        warnings=warnings,
    )


def parse_cv_import(content: bytes, filename: str, ext: str) -> CvImportProposal:
    if ext == "txt":
        text = content.decode("utf-8")
        if len(text) > MAX_EXTRACTED_CHARS:
            raise CvParserRejected("Extracted text limit exceeded")
        warnings = []
    else:
        parsed = parse_cv(content, filename, ext)
        text, warnings = parsed.extracted_text, parsed.warnings
    return _structure_text(text, filename, warnings)


_HEADINGS = {
    "summary": "summary", "profile": "summary", "experience": "experience",
    "work experience": "experience", "employment": "experience",
    "achievements": "achievements", "skills": "skills", "education": "education",
    "projects": "projects", "certifications": "certifications",
}


def _structure_text(text: str, filename: str, warnings: list[str]) -> CvImportProposal:
    grouped: list[tuple[str, str, list[str]]] = []
    kind, title, lines = "summary", "Summary", []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = _HEADINGS.get(re.sub(r"[:\s]+$", "", line.lower()))
        if heading:
            if lines:
                grouped.append((kind, title, lines))
            kind, title, lines = heading, line.rstrip(":"), []
        else:
            lines.append(line)
    if lines:
        grouped.append((kind, title, lines))
    sections = []
    for section_position, (section_kind, section_title, entries) in enumerate(grouped):
        import_entries = []
        for entry_position, body in enumerate(entries):
            claim_kind = _claim_kind(section_kind, body)
            import_entries.append({
                "id": f"entry-{section_position}-{entry_position}", "body": body,
                "position": entry_position,
                "claim": None if claim_kind is None else {
                    "kind": claim_kind, "content": {"statement": body},
                    "provenance": "imported",
                },
            })
        sections.append({
            "id": f"section-{section_position}-{section_kind}", "kind": section_kind,
            "title": section_title, "visible": True, "position": section_position,
            "entries": import_entries,
        })
    proposal_warnings = list(warnings)
    if not sections:
        proposal_warnings.append("No reviewable sections were detected.")
    return CvImportProposal(
        filename=filename, import_id=str(uuid.uuid4()),
        name=Path(filename).stem or "Imported CV",
        sections=sections, warnings=proposal_warnings,
    )


def _claim_kind(section_kind: str, body: str) -> str | None:
    if section_kind == "achievements" or re.search(r"\b\d+(?:[.%]|\b)", body):
        return "achievement"
    return {
        "skills": "skill", "education": "education", "projects": "project",
        "certifications": "certification", "experience": "experience",
    }.get(section_kind)


def _extract_pdf(content: bytes) -> str:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        if doc.needs_pass:
            raise CvParserRejected("Encrypted PDFs are not supported")
        if doc.page_count > MAX_PDF_PAGES:
            raise CvParserRejected("PDF page limit exceeded")
        pages = _bounded_text_parts(page.get_text() for page in doc)
    finally:
        doc.close()
    return "\n".join(pages)


def _extract_docx(content: bytes) -> str:
    import io

    from docx import Document
    from docx.table import Table

    with io.BytesIO(content) as buffer:
        doc = Document(buffer)
        text_parts = []
        blocks = list(doc.iter_inner_content())
        if not blocks:
            blocks = list(doc.paragraphs)
        for block in blocks:
            if isinstance(block, Table):
                text_parts.extend(
                    paragraph.text
                    for row in block.rows
                    for cell in row.cells
                    for paragraph in cell.paragraphs
                )
            else:
                text_parts.append(block.text)
        paragraphs = _bounded_text_parts(text_parts)
    return "\n".join(paragraphs)


def _bounded_text_parts(parts) -> list[str]:
    bounded: list[str] = []
    chars_count = 0
    for part in parts:
        chars_count += len(part)
        if bounded:
            chars_count += 1
        if chars_count > MAX_EXTRACTED_CHARS:
            raise CvParserRejected("Extracted text limit exceeded")
        bounded.append(part)
    return bounded
