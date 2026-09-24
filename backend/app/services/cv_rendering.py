"""Canonical CV rendering and deterministic artifact generation.

The JSON render model is the single source consumed by browser preview and both
exporters. PDF bytes are stable through ReportLab's invariant mode. DOCX semantic
content is deterministic; ZIP member timestamps/order are canonicalized so bytes
are stable with a fixed python-docx version.

Backward compatibility contract: ``build_render_model(document, template_id)``
called without a ``style`` argument must reproduce byte-identical PDF/DOCX output
for the three original templates (ats-essential, professional-editorial,
technical-portfolio) against entries that carry only ``body`` text — this is
covered by the golden pixel/layout snapshots in ``tests/test_cv_rendering.py``.
All style/structured-entry behavior below is additive and only engages when a
``style`` is supplied or an entry carries structured fields.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from html import escape

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    FrameBreak,
    KeepTogether,
    PageTemplate,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.schemas.cv_documents import CvArtifactEvidence, CvRenderModel, CvStyle
from app.services.cv_fonts import docx_font_name, pdf_font_names
from app.services.cv_parser import parse_cv_import

URL_RE = re.compile(r"https?://[^\s<>()\[\]{}\"']*[^\s<>()\[\]{}\"'.,;:!?]")


@dataclass(frozen=True)
class Template:
    id: str
    font: str
    accent: str
    body_size: int
    heading_size: int
    margin_mm: int
    section_gap: int
    align: str


TEMPLATES = {
    "ats-essential": Template("ats-essential", "Helvetica", "#111827", 10, 13, 18, 6, "left"),
    "professional-editorial": Template(
        "professional-editorial", "Times-Roman", "#7C2D12", 10, 15, 20, 8, "left"
    ),
    "technical-portfolio": Template(
        "technical-portfolio", "Courier", "#075985", 9, 12, 16, 7, "left"
    ),
    "modern-two-column": Template(
        "modern-two-column", "Helvetica", "#1D4ED8", 9, 12, 14, 6, "left"
    ),
    "minimal-serif": Template("minimal-serif", "Times-Roman", "#374151", 10, 13, 20, 7, "left"),
}

# Templates safe for strict ATS parsers without forcing ats_mode: single column,
# no sidebar, no decorative structure that could scramble reading order.
ATS_SAFE_TEMPLATES = {"ats-essential", "professional-editorial", "technical-portfolio", "minimal-serif"}

DENSITY_SCALE = {"compact": 0.88, "normal": 1.0, "spacious": 1.15}
DENSITY_GAP_SCALE = {"compact": 0.7, "normal": 1.0, "spacious": 1.4}

_BUILTIN_BOLD = {"Helvetica": "Helvetica-Bold", "Times-Roman": "Times-Bold", "Courier": "Courier-Bold"}


@dataclass(frozen=True)
class EffectiveStyle:
    layout_template_id: str
    font_name: str
    font_bold_name: str
    font_docx_name: str
    accent: str
    density: str
    ats_mode: bool
    body_size: int
    heading_size: int
    margin_mm: int
    section_gap: int
    align: str
    two_column: bool


def resolve_effective_style(template_id: str, style: CvStyle | None) -> EffectiveStyle:
    """Resolve template + style into the concrete tokens both renderers consume.

    ``style is None`` is the legacy path: it must reproduce the exact per-template
    defaults that existed before style customization shipped (no density/ats_mode
    behavior at all), so unmodified callers keep byte-identical output.
    """
    if style is None:
        t = TEMPLATES[template_id]
        return EffectiveStyle(
            layout_template_id=template_id,
            font_name=t.font,
            font_bold_name=_BUILTIN_BOLD.get(t.font, t.font),
            font_docx_name=t.font,
            accent=t.accent,
            density="normal",
            ats_mode=False,
            body_size=t.body_size,
            heading_size=t.heading_size,
            margin_mm=t.margin_mm,
            section_gap=t.section_gap,
            align=t.align,
            two_column=template_id == "modern-two-column",
        )
    ats = style.ats_mode
    layout_id = "ats-essential" if ats else template_id
    t = TEMPLATES[layout_id]
    if ats:
        font_name, font_bold, font_docx, accent = "Helvetica", "Helvetica-Bold", "Helvetica", "#111827"
    else:
        font_name, font_bold = pdf_font_names(style.font_id)
        font_docx = docx_font_name(style.font_id)
        accent = style.accent_color
    density = "normal" if ats else style.density
    scale = DENSITY_SCALE[density]
    gap_scale = DENSITY_GAP_SCALE[density]
    return EffectiveStyle(
        layout_template_id=layout_id,
        font_name=font_name,
        font_bold_name=font_bold,
        font_docx_name=font_docx,
        accent=accent,
        density=density,
        ats_mode=ats,
        body_size=max(8, round(t.body_size * scale)),
        heading_size=max(10, round(t.heading_size * scale)),
        margin_mm=t.margin_mm,
        section_gap=max(3, round(t.section_gap * gap_scale)),
        align=t.align,
        two_column=(not ats) and layout_id == "modern-two-column",
    )


def _clean(value) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def build_render_model(document, template_id: str, style: CvStyle | None = None) -> CvRenderModel:
    effective = resolve_effective_style(template_id, style)
    sections = []
    for section in sorted(document.sections, key=lambda item: (item["position"], item["id"])):
        if not section.get("visible", True):
            continue
        entries = []
        for entry in sorted(section["entries"], key=lambda item: (item["position"], item["id"])):
            text = " ".join(str(entry["body"]).split())
            bullets = [cleaned for b in entry.get("bullets") or [] if (cleaned := _clean(b))]
            link_source = " ".join([text, *bullets]) if bullets else text
            entries.append(
                {
                    "id": entry["id"],
                    "text": text,
                    "links": URL_RE.findall(link_source),
                    "heading": _clean(entry.get("heading")),
                    "subheading": _clean(entry.get("subheading")),
                    "location": _clean(entry.get("location")),
                    "start_date": _clean(entry.get("start_date")),
                    "end_date": _clean(entry.get("end_date")),
                    "bullets": bullets,
                }
            )
        sections.append(
            {
                "id": section["id"],
                "kind": section["kind"],
                "title": section["title"].strip(),
                "entries": entries,
            }
        )
    canonical = {
        "document_id": document.id,
        "document_name": document.name.strip(),
        "template_id": effective.layout_template_id,
        "page": {"width_mm": 210, "height_mm": 297, "margin_mm": effective.margin_mm},
        "tokens": {
            "font": effective.font_name,
            "font_bold": effective.font_bold_name,
            "font_docx": effective.font_docx_name,
            "accent": effective.accent,
            "body_size_pt": effective.body_size,
            "heading_size_pt": effective.heading_size,
            "section_gap_pt": effective.section_gap,
            "align": effective.align,
            "density": effective.density,
            "ats_mode": effective.ats_mode,
            "two_column": effective.two_column,
        },
        "sections": sections,
    }
    digest = hashlib.sha256(
        json.dumps(canonical, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    return CvRenderModel(**canonical, canonical_hash=digest)


def _markup(text: str) -> str:
    cursor, parts = 0, []
    for match in URL_RE.finditer(text):
        parts.extend(
            (
                escape(text[cursor : match.start()]),
                f'<a href="{escape(match.group())}" color="blue">{escape(match.group())}</a>',
            )
        )
        cursor = match.end()
    parts.append(escape(text[cursor:]))
    return "".join(parts)


def _pdf_styles(model: CvRenderModel) -> dict[str, ParagraphStyle]:
    tokens = model.tokens
    font = str(tokens["font"])
    font_bold = str(tokens.get("font_bold", font))
    accent = str(tokens["accent"])
    body_size = int(tokens["body_size_pt"])
    heading_size = int(tokens["heading_size_pt"])
    heading = ParagraphStyle(
        "heading",
        fontName=font,
        fontSize=heading_size,
        leading=heading_size + 3,
        textColor=HexColor(accent),
        spaceAfter=3,
        keepWithNext=True,
    )
    body = ParagraphStyle(
        "body", fontName=font, fontSize=body_size, leading=body_size + 3, spaceAfter=4
    )
    title = ParagraphStyle(
        "title",
        parent=heading,
        fontSize=heading_size + 6,
        leading=heading_size + 9,
        alignment=TA_CENTER if model.template_id == "professional-editorial" else TA_LEFT,
        spaceAfter=10,
    )
    entry_heading = ParagraphStyle(
        "entry_heading",
        fontName=font_bold,
        fontSize=body_size + 1,
        leading=body_size + 4,
        spaceAfter=1,
    )
    entry_heading_right = ParagraphStyle(
        "entry_heading_right", parent=entry_heading, alignment=TA_RIGHT
    )
    entry_meta = ParagraphStyle(
        "entry_meta",
        fontName=font,
        fontSize=max(7, body_size - 1),
        leading=body_size + 2,
        textColor=HexColor("#4B5563"),
        spaceAfter=2,
    )
    bullet = ParagraphStyle(
        "bullet",
        parent=body,
        leftIndent=10,
        bulletIndent=0,
        spaceAfter=2,
    )
    return {
        "heading": heading,
        "body": body,
        "title": title,
        "entry_heading": entry_heading,
        "entry_heading_right": entry_heading_right,
        "entry_meta": entry_meta,
        "bullet": bullet,
    }


def entry_render_lines(entry) -> list[str]:
    """The literal text lines an entry renders as, in rendering order.

    Single source of truth for what a structured entry (heading/subheading/
    location/dates/bullets) actually puts on the page: ``_entry_flow`` and
    ``_add_docx_structured_entry`` render exactly these lines (with their own
    formatting), and ``validate_artifact`` checks against exactly these lines
    instead of the unrendered ``entry.text`` — so a structured entry that
    renders as "heading + bullets" is no longer judged against its plain body
    (#322). A legacy/freeform entry (no heading) returns ``[entry.text]``.
    """
    if entry.heading is None:
        return [entry.text] if entry.text else []
    head_line = entry.heading
    if entry.subheading:
        head_line += f" — {entry.subheading}"
    date_range = " – ".join(part for part in (entry.start_date, entry.end_date) if part)
    if date_range:
        head_line += f" {date_range}"
    lines = [head_line]
    if entry.location:
        lines.append(entry.location)
    if entry.bullets:
        lines.extend(entry.bullets)
    elif entry.text and entry.text != entry.heading:
        lines.append(entry.text)
    return lines


def _entry_flow(entry, styles: dict[str, ParagraphStyle], content_width: float) -> list:
    if entry.heading is None:
        # Legacy/freeform path — must stay byte-identical to the pre-style renderer.
        return [Paragraph(_markup(entry.text), styles["body"])]
    flow: list = []
    left = f"<b>{escape(entry.heading)}</b>"
    if entry.subheading:
        left += f" — {escape(entry.subheading)}"
    date_range = " – ".join(part for part in (entry.start_date, entry.end_date) if part)
    if date_range:
        row = Table(
            [
                [
                    Paragraph(left, styles["entry_heading"]),
                    Paragraph(escape(date_range), styles["entry_heading_right"]),
                ]
            ],
            colWidths=[content_width * 0.7, content_width * 0.3],
        )
        row.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
        flow.append(row)
    else:
        flow.append(Paragraph(left, styles["entry_heading"]))
    if entry.location:
        flow.append(Paragraph(escape(entry.location), styles["entry_meta"]))
    for bullet_text in entry.bullets:
        flow.append(Paragraph(f"&bull;&nbsp;{_markup(bullet_text)}", styles["bullet"]))
    if not entry.bullets and entry.text and entry.text != entry.heading:
        flow.append(Paragraph(_markup(entry.text), styles["body"]))
    return flow


def render_pdf(model: CvRenderModel) -> bytes:
    tokens = model.tokens
    margin_mm = int(model.page["margin_mm"])
    section_gap = int(tokens["section_gap_pt"])
    styles = _pdf_styles(model)
    if bool(tokens.get("two_column", False)):
        return _render_pdf_two_column(model, styles, margin_mm, section_gap)
    return _render_pdf_single_column(model, styles, margin_mm, section_gap)


def _render_pdf_single_column(model, styles, margin_mm, section_gap) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=margin_mm * mm,
        rightMargin=margin_mm * mm,
        topMargin=margin_mm * mm,
        bottomMargin=margin_mm * mm,
        invariant=1,
        title=model.document_name,
        author="Career Workbench",
        creator="Career Workbench",
    )
    content_width = A4[0] - 2 * margin_mm * mm
    story = [Paragraph(escape(model.document_name), styles["title"])]
    for section in model.sections:
        flow = [Paragraph(escape(section.title), styles["heading"])]
        for entry in section.entries:
            flow.extend(_entry_flow(entry, styles, content_width))
        story.extend([KeepTogether(flow), Spacer(1, section_gap)])
    doc.build(story)
    return out.getvalue()


def _render_pdf_two_column(model, styles, margin_mm, section_gap) -> bytes:
    out = io.BytesIO()
    page_w, page_h = A4
    m = margin_mm * mm
    sidebar_w = 58 * mm
    gap = 6 * mm
    main_x = m + sidebar_w + gap
    side_frame = Frame(
        m, m, sidebar_w, page_h - 2 * m, id="side", leftPadding=0, rightPadding=6, topPadding=0
    )
    main_frame = Frame(
        main_x, m, page_w - main_x - m, page_h - 2 * m, id="main", leftPadding=0, topPadding=0
    )
    doc = BaseDocTemplate(
        out,
        pagesize=A4,
        leftMargin=m,
        rightMargin=m,
        topMargin=m,
        bottomMargin=m,
        invariant=1,
        title=model.document_name,
        author="Career Workbench",
        creator="Career Workbench",
    )
    doc.addPageTemplates([PageTemplate(id="two-col", frames=[side_frame, main_frame])])
    sidebar_kinds = {"skills", "certifications"}
    side_sections = [s for s in model.sections if s.kind in sidebar_kinds]
    main_sections = [s for s in model.sections if s.kind not in sidebar_kinds]
    if not side_sections and main_sections:
        side_sections, main_sections = main_sections[:1], main_sections[1:]
    story: list = [Paragraph(escape(model.document_name), styles["title"])]
    for section in side_sections:
        flow = [Paragraph(escape(section.title), styles["heading"])]
        for entry in section.entries:
            flow.extend(_entry_flow(entry, styles, sidebar_w))
        story.extend([KeepTogether(flow), Spacer(1, section_gap)])
    story.append(FrameBreak())
    for section in main_sections:
        flow = [Paragraph(escape(section.title), styles["heading"])]
        for entry in section.entries:
            flow.extend(_entry_flow(entry, styles, page_w - main_x - m))
        story.extend([KeepTogether(flow), Spacer(1, section_gap)])
    doc.build(story)
    return out.getvalue()


def render_docx(model: CvRenderModel) -> bytes:
    tokens = model.tokens
    font = str(tokens.get("font_docx", tokens["font"]))
    accent = str(tokens["accent"])
    body_size = int(tokens["body_size_pt"])
    heading_size = int(tokens["heading_size_pt"])
    margin_mm = int(model.page["margin_mm"])
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = (
        Inches(margin_mm / 25.4)
    )
    core = doc.core_properties
    core.title = model.document_name
    core.author = "Career Workbench"
    core.created = core.modified = datetime(2000, 1, 1, tzinfo=UTC)
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(body_size)
    title = doc.add_paragraph()
    title.alignment = 1 if model.template_id == "professional-editorial" else 0
    run = title.add_run(model.document_name)
    run.bold = True
    run.font.name = font
    run.font.size = Pt(heading_size + 6)
    run.font.color.rgb = RGBColor.from_string(accent[1:])
    content_width_in = (
        8.27 - 2 * margin_mm / 25.4
    )  # A4 width in inches minus margins, for the date tab stop
    for rendered_section in model.sections:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        r = p.add_run(rendered_section.title)
        r.bold = True
        r.font.name = font
        r.font.size = Pt(heading_size)
        r.font.color.rgb = RGBColor.from_string(accent[1:])
        for entry in rendered_section.entries:
            if entry.heading is None:
                p = doc.add_paragraph()
                p.paragraph_format.keep_together = True
                cursor = 0
                for match in URL_RE.finditer(entry.text):
                    p.add_run(entry.text[cursor : match.start()])
                    _add_hyperlink(p, match.group())
                    cursor = match.end()
                p.add_run(entry.text[cursor:])
                continue
            _add_docx_structured_entry(doc, entry, font, content_width_in)
    raw = io.BytesIO()
    doc.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    final = io.BytesIO()
    with zipfile.ZipFile(final, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    source.close()
    return final.getvalue()


def _add_docx_structured_entry(doc, entry, font: str, content_width_in: float) -> None:
    heading_line = doc.add_paragraph()
    heading_line.paragraph_format.keep_with_next = True
    date_range = " – ".join(part for part in (entry.start_date, entry.end_date) if part)
    if date_range:
        heading_line.paragraph_format.tab_stops.add_tab_stop(
            Inches(content_width_in), WD_TAB_ALIGNMENT.RIGHT
        )
    heading_run = heading_line.add_run(entry.heading)
    heading_run.bold = True
    heading_run.font.name = font
    if entry.subheading:
        sub_run = heading_line.add_run(f" — {entry.subheading}")
        sub_run.font.name = font
    if date_range:
        date_run = heading_line.add_run(f"\t{date_range}")
        date_run.font.name = font
    if entry.location:
        location_p = doc.add_paragraph()
        location_run = location_p.add_run(entry.location)
        location_run.italic = True
        location_run.font.name = font
    for bullet_text in entry.bullets:
        bullet_p = doc.add_paragraph(style="List Bullet")
        bullet_run = bullet_p.add_run(bullet_text)
        bullet_run.font.name = font
    if not entry.bullets and entry.text and entry.text != entry.heading:
        body_p = doc.add_paragraph()
        body_run = body_p.add_run(entry.text)
        body_run.font.name = font


def _add_hyperlink(paragraph, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    text = OxmlElement("w:t")
    text.text = url
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


_BULLET_MARKER_RE = re.compile(r"^[•◦\-*]\s*")


def _normalize_text(value: str) -> str:
    return " ".join(str(value).split())


def _strip_bullet_marker(value: str) -> str:
    """Drop a leading bullet glyph a PDF's extracted text carries as literal
    characters (rendered via ``&bull;&nbsp;`` in ``_entry_flow``) so it never
    changes the *content* comparison in ``validate_artifact`` — DOCX bullets
    use paragraph-level list formatting and never carry the glyph at all."""
    return _BULLET_MARKER_RE.sub("", value)


def validate_artifact(model: CvRenderModel, artifact: bytes, fmt: str) -> CvArtifactEvidence:
    parsed = parse_cv_import(artifact, f"cv.{fmt}", fmt)
    # A structured entry renders as several lines (heading, location, one per
    # bullet); own-parser re-import produces one entry per rendered line, so
    # comparing per-section joined text (rather than an exact per-entry list)
    # is what stays stable across both freeform and structured entries (#322).
    expected_structure = [
        (
            section.kind,
            section.title,
            _normalize_text(
                " ".join(line for entry in section.entries for line in entry_render_lines(entry))
            ),
        )
        for section in model.sections
    ]
    actual_structure = [
        (
            section.kind,
            section.title,
            _normalize_text(
                " ".join(_strip_bullet_marker(entry.body) for entry in section.entries)
            ),
        )
        for section in parsed.sections
    ]
    if actual_structure and actual_structure[0][2] == _normalize_text(model.document_name):
        actual_structure = actual_structure[1:]
    content_equivalent = expected_structure == actual_structure
    links = [
        link for section in model.sections for entry in section.entries for link in entry.links
    ]
    extracted = "\n".join(e.body for s in parsed.sections for e in s.entries)
    if fmt == "pdf":
        import fitz

        with fitz.open(stream=artifact, filetype="pdf") as rendered:
            page_lines = [page.get_text().splitlines() for page in rendered]
            page_text = [" ".join(" ".join(lines).split()) for lines in page_lines]
            page_breaks_ok = rendered.page_count > 0 and all(page_lines)
            for section in model.sections:
                if not section.entries:
                    continue
                # Checked per rendered line (not one joined block): a bullet
                # line's literal "•" glyph in the extracted text would
                # otherwise break a contiguous substring match even though
                # KeepTogether already guarantees the whole entry shares a
                # page with its heading (#322).
                first_lines = [
                    _normalize_text(line) for line in entry_render_lines(section.entries[0])
                ]
                page_breaks_ok = page_breaks_ok and any(
                    section.title in text and all(line in text for line in first_lines)
                    for text in page_text
                )
            artifact_links = {link.get("uri") for page in rendered for link in page.get_links()}
    else:
        with zipfile.ZipFile(io.BytesIO(artifact)) as package:
            document_xml = package.read("word/document.xml")
            relationships = package.read("word/_rels/document.xml.rels")
        page_breaks_ok = b"w:keepNext" in document_xml and b"w:sectPr" in document_xml
        artifact_links = {link for link in links if link.encode() in relationships}
    return CvArtifactEvidence(
        template_id=model.template_id,
        format=fmt,
        searchable_text="pass" if content_equivalent else "fail",
        links="pass"
        if all(link in extracted and link in artifact_links for link in links)
        else "fail",
        page_breaks="pass" if page_breaks_ok else "fail",
        re_importability="pass" if content_equivalent else "fail",
        canonical_hash=model.canonical_hash,
    )
